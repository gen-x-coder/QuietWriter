from pathlib import Path

import pytest
from docx import Document

import quietwriter.storage as storage_module
from quietwriter.storage import Library, StorageWriteError


def _docx(path: Path, *, title: str = 'Transactioneel boek') -> Path:
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading('Hoofdstuk Een', level=1)
    doc.add_paragraph('Eerste alinea.')
    doc.add_heading('Hoofdstuk Twee', level=1)
    doc.add_paragraph('Tweede alinea.')
    doc.save(path)
    return path


def _visible_book_dirs(lib: Library) -> list[Path]:
    return sorted(p for p in lib.books_dir.iterdir() if p.is_dir() and not p.name.startswith('.'))


def _staging_dirs(lib: Library) -> list[Path]:
    return sorted(p for p in lib.books_dir.iterdir() if p.is_dir() and p.name.startswith('.import-'))


def test_docx_import_is_published_only_after_complete_stage(tmp_path: Path):
    source = _docx(tmp_path / 'book.docx')
    lib = Library(tmp_path / 'workspace')

    book, warnings = lib.import_docx_book(source)

    assert warnings == ()
    assert book.path.is_dir()
    assert not book.path.name.startswith('.')
    assert _staging_dirs(lib) == []
    assert _visible_book_dirs(lib) == [book.path]
    reloaded = lib.load_book(book.path)
    assert [chapter.title for section in reloaded.sections for chapter in section.chapters] == [
        'Hoofdstuk Een', 'Hoofdstuk Twee'
    ]
    assert all((book.path / chapter.file).is_file() for section in reloaded.sections for chapter in section.chapters)


def test_docx_import_write_failure_leaves_no_half_book(tmp_path: Path, monkeypatch):
    source = _docx(tmp_path / 'book.docx')
    lib = Library(tmp_path / 'workspace')
    real_write = storage_module._safe_atomic_write_text
    chapter_writes = 0

    def fail_on_second_chapter(path, text, *args, **kwargs):
        nonlocal chapter_writes
        if Path(path).suffix == '.md' and 'chapters' in Path(path).parts:
            chapter_writes += 1
            if chapter_writes == 2:
                raise StorageWriteError(Path(path), OSError('gesimuleerde schijffout'))
        return real_write(path, text, *args, **kwargs)

    monkeypatch.setattr(storage_module, '_safe_atomic_write_text', fail_on_second_chapter)

    with pytest.raises(StorageWriteError):
        lib.import_docx_book(source)

    assert _visible_book_dirs(lib) == []
    assert _staging_dirs(lib) == []
    assert lib.list_books() == []


def test_docx_import_publish_failure_cleans_stage_and_stays_invisible(tmp_path: Path, monkeypatch):
    source = _docx(tmp_path / 'book.docx')
    lib = Library(tmp_path / 'workspace')
    real_replace = storage_module.os.replace

    def fail_directory_publish(src, dst):
        if Path(src).name.startswith('.import-'):
            raise PermissionError('gesimuleerde sync-lock')
        return real_replace(src, dst)

    monkeypatch.setattr(storage_module.os, 'replace', fail_directory_publish)

    with pytest.raises(PermissionError):
        lib.import_docx_book(source)

    assert _visible_book_dirs(lib) == []
    assert _staging_dirs(lib) == []
    assert lib.list_books() == []
