"""1.7: scènes in de tekst, personages herkennen, exportdoel Inzenden.

Wat hier NIET getest kan worden: hoe de regel onder de tekst, de "Volgende in
Planning"-hulp en het manuscript in Word er op Windows (thema's, DPI) uitzien.
"""
from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path

import pytest

from quietwriter.character_mentions import ChapterTextCache, chapter_mentions, character_terms, mention_counts, plain_text
from quietwriter.planning_models import Character, Scene
from quietwriter.scene_segments import chapter_segments, map_scenes, scenes_for_chapter, segment_index_for_line

ROOT = Path(__file__).resolve().parents[2]

TEXT = 'Anna liep **snel** weg.\n\n***\n\nBram bleef staan.\n***\n'


# --- segmenten ----------------------------------------------------------------------

def test_segments_follow_scene_breaks_and_keep_empty_last_segment():
    segs = chapter_segments(TEXT)
    assert [s.index for s in segs] == [0, 1, 2]
    assert segs[0].excerpt == 'Anna liep snel weg.' and segs[0].has_text
    assert segs[1].excerpt == 'Bram bleef staan.' and segs[1].first_line == 3
    assert not segs[2].has_text
    lines = TEXT.split('\n')
    assert segment_index_for_line(lines, 0) == 0
    assert segment_index_for_line(lines, 4) == 1
    assert segment_index_for_line(lines, 6) == 2


def test_empty_chapter_is_one_empty_segment_and_long_excerpt_is_cut():
    assert [(s.index, s.has_text) for s in chapter_segments('')] == [(0, False)]
    long = ' '.join(f'w{i}' for i in range(20))
    assert chapter_segments(long)[0].excerpt.endswith('…')


def test_map_scenes_is_positional_and_reports_unplanned_text():
    a, b = Scene(chapter_id='c1', title='A'), Scene(chapter_id='c1', title='B')
    other = Scene(chapter_id='c2', title='X')
    planned = scenes_for_chapter([a, other, b], 'c1')
    assert planned == [a, b]
    mapping = map_scenes([a], TEXT)
    assert mapping.scene_for_segment(0) is a and mapping.scene_for_segment(1) is None
    assert [s.excerpt for s in mapping.unplanned] == ['Bram bleef staan.']   # leeg laatste stuk telt niet
    assert map_scenes([a, b], TEXT).unplanned == ()
    assert mapping.segment_for_scene(a.id).index == 0


def test_open_points_and_markup_never_leak_into_excerpts():
    source = 'Hij zei <!--qw:todo:abcdef12-->iets<!--/qw:todo--> *zacht*.'
    assert chapter_segments(source)[0].excerpt == 'Hij zei iets zacht.'


# --- personages ---------------------------------------------------------------------

def _anna():
    return Character(id='a', name='Anna de Vries', aliases='An, mevrouw De Vries')


def test_terms_are_full_name_first_name_and_aliases():
    assert character_terms(_anna()) == ('Anna de Vries', 'Anna', 'An', 'mevrouw De Vries')
    assert character_terms(Character(id='b', name='Jo')) == ('Jo',)          # te korte voornaam: alleen volledige naam
    assert character_terms(Character(id='c', name='Roos')) == ('Roos',)


def test_mentions_are_case_sensitive_whole_words():
    roos = Character(id='r', name='Roos')
    text = plain_text('De roos bloeide. Roos lachte. Roosje niet. Roos’ jas. Roos\' hoed.')
    assert mention_counts(text, [roos]) == {'r': 3}
    anna = _anna()
    assert mention_counts('Anna de Vries kwam. Anna zei niets. An knikte. Annabel niet.', [anna]) == {'a': 3}


def test_first_appearance_is_relative_to_earlier_chapters():
    anna, bram = _anna(), Character(id='b', name='Bram')
    texts = [('c1', 'Anna was er.'), ('c2', 'Anna en Bram. Bram!'), ('c3', 'Niemand.')]
    rows = chapter_mentions('c2', texts, [anna, bram])
    assert [(r.name, r.count, r.first_here) for r in rows] == [('Bram', 2, True), ('Anna de Vries', 1, False)]
    assert chapter_mentions('c3', texts, [anna, bram]) == ()


def test_text_cache_rereads_only_after_change():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / 'c.md'
        path.write_text('Anna', encoding='utf-8')
        cache = ChapterTextCache()
        assert cache.text(path) == 'Anna'
        os.utime(path, ns=(1, 1))
        path.write_text('Bram is langer', encoding='utf-8')
        assert cache.text(path) == 'Bram is langer'
        assert cache.text(Path(td) / 'weg.md') is None              # ontbreekt is niet leeg (review 1.7.0)


# --- Planning-gegevens: bijnamen en vooruit-compatibiliteit -------------------------

def test_aliases_round_trip_and_old_files_still_load():
    c = Character.from_dict({'id': 'x', 'name': 'Anna', 'aliases': 'An'})
    assert c.aliases == 'An' and c.to_dict()['aliases'] == 'An'
    assert Character.from_dict({'id': 'y', 'name': 'Bram'}).aliases == ''


def test_unknown_fields_from_a_newer_version_survive_a_save():
    row = {'id': 'x', 'name': 'Anna', 'eye_color': 'groen', 'nested': {'a': 1}}
    assert Character.from_dict(row).to_dict()['eye_color'] == 'groen'
    scene = Scene.from_dict({'id': 's', 'title': 'T', 'anchor': {'v': 2}})
    assert scene.to_dict()['anchor'] == {'v': 2}
    assert 'extra' not in scene.to_dict()
    # Een bekend veld wint altijd van een gelijknamige 'extra'.
    assert Scene.from_dict({'id': 's', 'title': 'Echt'}).to_dict()['title'] == 'Echt'


def test_planning_store_keeps_unknown_fields_on_disk():
    from quietwriter.planning_storage import PlanningStore
    from quietwriter.storage import Library
    with tempfile.TemporaryDirectory() as td:
        lib = Library(Path(td) / 'ws'); book = lib.create_book('Pip'); lib.track_book(book)
        store = PlanningStore(lib)
        root = store.root(book); root.mkdir(parents=True, exist_ok=True)
        (root / 'characters.json').write_text(json.dumps({'version': 1, 'characters': [
            {'id': 'a', 'name': 'Anna', 'future': [1, 2]}]}), encoding='utf-8')
        lib.refresh_book_revision(book)
        chars = store.load_characters(book)
        chars[0].aliases = 'An'
        store.save_characters(book, chars)
        saved = json.loads((root / 'characters.json').read_text(encoding='utf-8'))['characters'][0]
        assert saved['future'] == [1, 2] and saved['aliases'] == 'An'


# --- Inzenden -----------------------------------------------------------------------

def _submission_book(td, author='Jan van der Berg', language='nl'):
    from quietwriter.storage import Library
    lib = Library(Path(td) / 'ws')
    book = lib.create_book('De Mist: een roman')
    book.metadata['author'] = author
    book.metadata['language'] = language
    lib.track_book(book)
    chapter = book.sections[0].chapters[0]
    lib.save_chapter(book, chapter, 'Eerste alinea met *cursief*.\nTweede alinea.\n\n***\n\nNa de breuk.\n> Citaat\nDaarna.')
    return lib, book


def _export(lib, book, td, options):
    from quietwriter.exporting import build_export_document
    from quietwriter.exporting.runner import destination_for, run_export
    from docx import Document
    document = build_export_document(lib, book)
    destination = destination_for(Path(td), document, 'submission')
    run_export(lib, book, document, 'submission', {'submission': options}, destination)
    return destination, Document(str(destination))


def test_submission_manuscript_format():
    from docx.enum.text import WD_LINE_SPACING
    with tempfile.TemporaryDirectory() as td:
        lib, book = _submission_book(td)
        destination, doc = _export(lib, book, td, {'contact': 'Jan\njan@example.org'})
        assert destination.name.endswith('-manuscript.docx')            # overschrijft de Word-export niet
        texts = [p.text for p in doc.paragraphs]
        assert texts[:2] == ['Jan', 'jan@example.org']
        assert 'DE MIST: EEN ROMAN' in texts and 'door Jan van der Berg' in texts
        assert any(t.startswith('ongeveer ') and t.endswith(' woorden') for t in texts)
        assert '#' in texts and texts[-1] == 'EINDE'
        body = {p.text: p for p in doc.paragraphs}
        assert body['Eerste alinea met cursief.'].paragraph_format.first_line_indent == 0   # eerste na titel
        assert body['Tweede alinea.'].paragraph_format.first_line_indent > 0
        assert body['Na de breuk.'].paragraph_format.first_line_indent == 0                 # eerste na scènebreuk
        assert body['Tweede alinea.'].paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
        header = ''.join(p.text for p in doc.sections[1].header.paragraphs)
        assert header.startswith('van der Berg / DE MIST / ')
        assert not ''.join(p.text for p in doc.sections[0].header.paragraphs)      # titelpagina zonder kop
        assert doc.styles['Normal'].font.name == 'Times New Roman'


def test_anonymous_submission_contains_no_name_anywhere():
    with tempfile.TemporaryDirectory() as td:
        lib, book = _submission_book(td)
        destination, doc = _export(lib, book, td, {'anonymous': True, 'contact': 'Jan\nStraat 1', 'font': 'arial',
                                                   'line_spacing': 'one_half'})
        everything = '\n'.join(p.text for p in doc.paragraphs)
        everything += '\n'.join(p.text for s in doc.sections for p in s.header.paragraphs)
        assert 'Berg' not in everything and 'Straat' not in everything
        assert doc.core_properties.author == '' and doc.core_properties.last_modified_by == ''
        assert doc.styles['Normal'].font.name == 'Arial'
        import zipfile
        with zipfile.ZipFile(destination) as archive:
            raw = b''.join(archive.read(name) for name in archive.namelist())
        assert b'Berg' not in raw


def test_submission_title_page_follows_book_language():
    with tempfile.TemporaryDirectory() as td:
        lib, book = _submission_book(td, author='Ann Lee', language='en')
        _, doc = _export(lib, book, td, {})
        texts = [p.text for p in doc.paragraphs]
        assert 'by Ann Lee' in texts and texts[-1] == 'THE END'


def test_rounded_word_count_and_surname_rules():
    from quietwriter.exporting.submission_exporter import format_number, rounded_word_count, surname
    assert rounded_word_count(87) == 90 and rounded_word_count(12_345) == 12_300 and rounded_word_count(82_499) == 82_000
    assert format_number(82_000, 'nl') == '82.000' and format_number(82_000, 'en') == '82,000'
    assert surname('Lucas Bonsel') == 'Bonsel' and surname('Jan van der Berg') == 'van der Berg' and surname('Plato') == 'Plato'


def test_export_settings_never_store_contact_details_in_the_book():
    from quietwriter.exporting.settings import ExportSettingsStore
    with tempfile.TemporaryDirectory() as td:
        lib, book = _submission_book(td)
        store = ExportSettingsStore(lib)
        store.save(book, {'format': 'submission', 'purpose': 'submission',
                          'submission': {'font': 'courier', 'contact': 'Geheim adres', 'anonymous': True}})
        raw = store.path(book).read_text(encoding='utf-8')
        assert 'Geheim' not in raw
        loaded = store.load(book)
        assert loaded['format'] == 'submission' and loaded['purpose'] == 'submission'
        assert loaded['submission']['font'] == 'courier' and loaded['submission']['anonymous'] is True


def test_preflight_mentions_images_for_submission():
    from quietwriter.exporting.preflight import run_preflight
    from quietwriter.exporting.models import ExportAsset, ExportChapter, ExportDocument, ExportSection
    doc = ExportDocument(id='b', title='T', author='A', language='nl', slug='t', metadata={},
                         sections=(ExportSection('s', '', (ExportChapter('c', 'H', 'x'),)),),
                         front_matter=(), back_matter=(),
                         assets=(ExportAsset('i', 'a.png', 'image/png', b'', 'inline'),))
    keys = [item.key for item in run_preflight(doc, 'submission', {}).items]
    assert 'submission_images' in keys
    assert run_preflight(doc, 'submission', {}).can_export


# --- in het echte venster ----------------------------------------------------------

@pytest.fixture(scope='module')
def app():
    pytest.importorskip('PySide6')
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _window(root):
    from PySide6.QtCore import QSettings
    from quietwriter.storage import Library
    from quietwriter.ui.main_window import MainWindow
    settings = QSettings(str(root / 'settings.ini'), QSettings.IniFormat)
    return MainWindow(settings, Library(root / 'workspace'), [])


@pytest.mark.qt
def test_editor_segments_jump_and_tail_ghost(app):
    from quietwriter.ui.manuscript_editor import ManuscriptEditor
    editor = ManuscriptEditor()
    editor.setPlainText('Eerste.\n***\nTweede.\n***\n')
    editor.set_planning_scenes([Scene(title='A'), Scene(title='B', status='geschreven'), Scene(title='C'), Scene(title='D')])
    assert len(editor.scene_break_blocks()) == 2
    assert editor.go_to_segment(1)
    assert editor.textCursor().block().text() == 'Tweede.' and editor.segment_at_cursor() == 1
    assert editor.go_to_segment(2) and editor.segment_at_cursor() == 2
    assert not editor.go_to_segment(5)
    block, scenes = editor._tail_ghost_scenes()
    assert block is not None and [s.title for s in scenes] == ['C', 'D']        # na de laatste breuk, open scènes
    before = editor.source_text()
    editor.set_scene_hint('Scène 3 in Planning: C')
    assert editor.source_text() == before and editor.scene_hint.text()
    editor.setPlainText('Eerste.\n***\nTweede.')
    assert editor._tail_ghost_scenes() == (None, ())                         # laatste stuk heeft tekst


@pytest.mark.qt
def test_chapter_context_numbers_scenes_and_adds_unplanned_text(app):
    from quietwriter.planning_storage import PlanningStore
    with tempfile.TemporaryDirectory() as td:
        window = _window(Path(td))
        try:
            lib = window.library
            book = lib.create_book('Pip')
            chapter = book.sections[0].chapters[0]
            lib.track_book(book)
            lib.save_chapter(book, chapter, 'Anna kwam binnen.\n***\nBram ging weg.\n***\nAnna bleef.')
            store = PlanningStore(lib)
            store.save_characters(book, [Character(id='a', name='Anna'), Character(id='b', name='Bram')])
            store.save_scenes(book, [Scene(id='s1', chapter_id=chapter.id, title='Aankomst', synopsis='Anna komt.')])
            window.open_book(book)
            page = window.editor_page
            page.open_chapter(chapter)
            page.show_chapter_context()
            panel = page.chapter_context
            assert [seg.excerpt for seg in panel.segment_map.unplanned] == ['Bram ging weg.', 'Anna bleef.']
            assert [(m.name, m.count) for m in panel.mentions] == [('Anna', 2), ('Bram', 1)]
            page.editor.go_to_segment(0)
            page._scene_position_changed()
            assert page.editor._scene_hint_full.startswith('Scène 1 in Planning: Aankomst — Anna komt.')
            page._add_unplanned_scenes()
            scenes = store.load_scenes(book)
            assert [(s.title, s.status) for s in scenes] == [
                ('Aankomst', 'idee'), ('Bram ging weg.', 'geschreven'), ('Anna bleef.', 'geschreven')]
            assert lib.read_chapter(book, chapter) == 'Anna kwam binnen.\n***\nBram ging weg.\n***\nAnna bleef.'
            assert panel.segment_map.unplanned == ()
            page.editor.go_to_segment(2)
            page._scene_position_changed()
            assert 'Scène 3 in Planning: Anna bleef.' in page.editor._scene_hint_full
        finally:
            window.search_index.close(); window.close(); app.processEvents()


@pytest.mark.qt
def test_wizard_submission_keeps_contact_on_this_computer(app, monkeypatch):
    with tempfile.TemporaryDirectory() as td:
        window = _window(Path(td))
        try:
            lib = window.library
            book = lib.create_book('Pip'); book.metadata['author'] = 'Ann Lee'
            lib.track_book(book)
            lib.save_chapter(book, book.sections[0].chapters[0], 'Tekst.')
            window.settings.setValue('export/submission_contact', 'Ann Lee\nann@example.org')
            window.open_book(book)
            wizard = window.export_page.wizard if hasattr(window, 'export_page') else None
            if wizard is None:
                pytest.skip('Exportpagina niet beschikbaar in deze opbouw')
            wizard.set_book(book, reset=True)
            wizard._set_purpose('submission')
            wizard.step = 'appearance'
            wizard._render_appearance()
            assert wizard.submission_card.isVisibleTo(wizard) or not wizard.isVisible()
            assert wizard.submission_contact.toPlainText() == 'Ann Lee\nann@example.org'
            wizard.submission_one_half.setChecked(True)
            wizard.submission_anonymous.setChecked(True)
            assert wizard.draft['submission']['line_spacing'] == 'one_half'
            assert wizard.draft['submission']['anonymous'] is True
            assert not wizard.submission_contact.isEnabled()
        finally:
            window.search_index.close(); window.close(); app.processEvents()


# --- ChatGPT-review 1.7.0 ----------------------------------------------------------

def test_shared_first_name_counts_for_nobody_but_full_names_still_work():
    a, b = Character(id='a', name='Roos van Dijk'), Character(id='b', name='Roos de Wit')
    assert mention_counts('Roos wandelde naar huis.', [a, b]) == {}
    assert mention_counts('Roos van Dijk zat. Roos de Wit ook.', [a, b]) == {'a': 1, 'b': 1}


def test_explicit_name_beats_another_characters_first_name():
    anna, bakker = Character(id='a', name='Anna'), Character(id='b', name='Anna Bakker')
    assert mention_counts('Anna kwam. Anna Bakker ging.', [anna, bakker]) == {'a': 1, 'b': 1}
    shared = [Character(id='x', name='Xavier', aliases='Dokter'), Character(id='y', name='Yara', aliases='Dokter')]
    assert mention_counts('De Dokter kwam.', shared) == {}


def test_missing_earlier_chapter_makes_first_appearance_unknown():
    anna = Character(id='a', name='Anna')
    rows = chapter_mentions('c2', [('c1', None), ('c2', 'Anna.')], [anna])
    assert rows[0].first_here is None
    rows = chapter_mentions('c2', [('c1', ''), ('c2', 'Anna.')], [anna])
    assert rows[0].first_here is True                                   # leeg is wél bekend


def test_cache_distinguishes_missing_and_unreadable_from_empty():
    with tempfile.TemporaryDirectory() as td:
        cache = ChapterTextCache()
        empty = Path(td) / 'leeg.md'; empty.write_text('', encoding='utf-8')
        broken = Path(td) / 'kapot.md'; broken.write_bytes(b'\xff\xfe\xfa')
        assert cache.text(empty) == ''
        assert cache.text(broken) is None
        assert cache.text(Path(td) / 'weg.md') is None


@pytest.mark.qt
def test_add_unplanned_is_one_fresh_action_for_panel_and_outline(app):
    from quietwriter.planning_storage import PlanningStore
    with tempfile.TemporaryDirectory() as td:
        window = _window(Path(td))
        try:
            lib = window.library
            book = lib.create_book('Pip')
            chapter = book.sections[0].chapters[0]
            lib.track_book(book)
            lib.save_chapter(book, chapter, 'Een.\n***\nTwee.')
            store = PlanningStore(lib)
            store.save_scenes(book, [])
            window.open_book(book)
            planning = window.planning_page
            outline = planning.outline_page
            outline.load()
            stale = list(outline.scenes)
            # Iemand anders (of het paneel) voegt intussen al scènes toe.
            assert planning.add_unplanned_scenes(chapter.id, 'Een.\n***\nTwee.') == ('mine', 2)
            outline.scenes = stale                                   # verouderd scherm
            outline.add_unplanned(chapter.id)                       # mag niets dubbel maken
            assert [s.title for s in store.load_scenes(book)] == ['Een.', 'Twee.']
            assert planning.add_unplanned_scenes(chapter.id, 'Een.\n***\nTwee.') == ('nothing', 0)
        finally:
            window.search_index.close(); window.close(); app.processEvents()


@pytest.mark.qt
def test_scene_hint_follows_a_renamed_scene_after_returning_from_planning(app):
    from quietwriter.planning_storage import PlanningStore
    with tempfile.TemporaryDirectory() as td:
        window = _window(Path(td))
        try:
            lib = window.library
            book = lib.create_book('Pip')
            chapter = book.sections[0].chapters[0]
            lib.track_book(book)
            lib.save_chapter(book, chapter, 'Een.')
            store = PlanningStore(lib)
            store.save_scenes(book, [Scene(id='s', chapter_id=chapter.id, title='Oud')])
            window.open_book(book)
            page = window.editor_page
            page.open_chapter(chapter)
            assert 'Oud' in page.editor._scene_hint_full
            scenes = store.load_scenes(book); scenes[0].title = 'Nieuw'; store.save_scenes(book, scenes)
            page.refresh_chapter_context()                         # zoals bij terugkeer uit Planning
            assert 'Nieuw' in page.editor._scene_hint_full
        finally:
            window.search_index.close(); window.close(); app.processEvents()
