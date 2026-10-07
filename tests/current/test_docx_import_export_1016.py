from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from quietwriter.docx_io import paragraph_to_markdown, parse_docx_book
from quietwriter.exporting.builder import _bcp47_language


def _add_hyperlink(paragraph, text: str, url='https://example.com'):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = text
    run.append(t); hyperlink.append(run); paragraph._p.append(hyperlink)


def test_page_oriented_import_keeps_long_first_paragraph_in_body(tmp_path: Path):
    path = tmp_path / 'long-first.docx'
    doc = Document()
    doc.add_paragraph('Titel', style='Title').add_run().add_break(WD_BREAK.PAGE)
    long_text = 'Het was een koude ochtend in maart toen Ilse de deur van de vuurtoren opende. ' * 4
    doc.add_paragraph(long_text).add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Volgende').add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)
    parsed = parse_docx_book(path)
    all_text = '\n'.join(ch['text'] for _, chapters in parsed.sections for ch in chapters)
    assert long_text.strip() in all_text
    assert any(ch['title'].startswith('Deel ') for _, chapters in parsed.sections for ch in chapters)


def test_hyperlink_text_is_preserved_in_import(tmp_path: Path):
    path = tmp_path / 'links.docx'
    doc = Document(); doc.add_heading('Hoofdstuk', level=1)
    p = doc.add_paragraph('Lees meer op '); _add_hyperlink(p, 'LINKTEKST_HIER'); p.add_run(' na de link.')
    doc.save(path)
    parsed = parse_docx_book(path)
    body = '\n'.join(ch['text'] for _, chapters in parsed.sections for ch in chapters)
    assert 'Lees meer op LINKTEKST_HIER na de link.' in body


def test_adjacent_equal_formatting_runs_are_merged_before_markdown():
    doc = Document(); p = doc.add_paragraph()
    for text in ('vet', ' gesplitst', ' woord'):
        r = p.add_run(text); r.bold = True
    assert paragraph_to_markdown(p) == '**vet gesplitst woord**'


def test_bcp47_language_normalization():
    assert _bcp47_language('nl_NL') == 'nl-NL'
    assert _bcp47_language('en_us') == 'en-US'
    assert _bcp47_language('nl') == 'nl'


def test_docx_import_escapes_literal_markdown_like_prose_instead_of_reinterpreting_it():
    from quietwriter.document_view import parse_document, visible_block_text

    doc = Document(); p = doc.add_paragraph('- Kom je? *316 5*3 **letterlijk**')
    source = paragraph_to_markdown(p)
    block = parse_document(source).blocks[0]
    assert block.kind == 'paragraph'
    assert block.inline_spans == ()
    assert visible_block_text(block) == '- Kom je? *316 5*3 **letterlijk**'


def test_docx_hard_break_keeps_formatting_local_to_each_quietwriter_paragraph():
    from quietwriter.document_view import parse_document, visible_block_text

    doc = Document(); p = doc.add_paragraph()
    run = p.add_run('vet voor'); run.bold = True
    run.add_break(); run.add_text('vet na')
    source = paragraph_to_markdown(p)
    assert source == '**vet voor\u2028vet na**'
    blocks = parse_document(source).blocks
    assert [visible_block_text(block) for block in blocks] == ['vet voor\u2028vet na']
    assert all(any(span.kind == 'bold' for span in block.inline_spans) for block in blocks)


def test_docx_character_styles_strong_and_emphasis_are_imported_semantically():
    doc = Document(); p = doc.add_paragraph()
    strong = p.add_run('sterk'); strong.style = 'Strong'
    p.add_run(' en ')
    emphasis = p.add_run('nadruk'); emphasis.style = 'Emphasis'
    assert paragraph_to_markdown(p) == '**sterk** en *nadruk*'


def test_docx_paragraph_style_character_formatting_is_inherited_by_runs():
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    style = doc.styles.add_style('Vet proza', WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    p = doc.add_paragraph(style='Vet proza'); p.add_run('hele alinea')
    assert paragraph_to_markdown(p) == '**hele alinea**'


def test_docx_direct_false_overrides_inherited_character_style():
    doc = Document(); p = doc.add_paragraph()
    run = p.add_run('niet vet'); run.style = 'Strong'; run.bold = False
    assert paragraph_to_markdown(p) == 'niet vet'


def test_docx_custom_paragraph_style_based_on_heading_is_recognized_as_structure(tmp_path: Path):
    from docx.enum.style import WD_STYLE_TYPE

    path = tmp_path / 'custom-heading.docx'
    doc = Document(); doc.core_properties.title = 'Roman'
    chapter_style = doc.styles.add_style('Mijn hoofdstuk', WD_STYLE_TYPE.PARAGRAPH)
    chapter_style.base_style = doc.styles['Heading 1']
    doc.add_paragraph('Eerste', style='Mijn hoofdstuk')
    doc.add_paragraph('Tekst')
    doc.add_paragraph('Tweede', style='Mijn hoofdstuk')
    doc.add_paragraph('Meer tekst')
    doc.save(path)

    parsed = parse_docx_book(path)
    chapters = parsed.sections[0][1]
    assert [c['title'] for c in chapters] == ['Eerste', 'Tweede']
    assert chapters[0]['text'] == 'Tekst'


def test_docx_hyperlink_import_warns_that_destination_is_not_preserved(tmp_path: Path):
    path = tmp_path / 'hyperlink-warning.docx'
    doc = Document(); doc.add_heading('Hoofdstuk', level=1)
    p = doc.add_paragraph('Zie '); _add_hyperlink(p, 'voorbeeld', 'https://example.com')
    doc.save(path)

    parsed = parse_docx_book(path)
    assert any('hyperlink' in warning.lower() for warning in parsed.warnings)


def test_page_oriented_short_first_sentence_is_kept_as_body_not_consumed_as_title(tmp_path: Path):
    path = tmp_path / 'short-first.docx'
    doc = Document()
    doc.add_paragraph('Titel', style='Title').add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Nee.')
    doc.add_paragraph('Daarna kwam de rest van de alinea.').add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Volgende', style='Title')
    doc.add_paragraph('Tekst')
    doc.save(path)

    parsed = parse_docx_book(path)
    bodies = '\n'.join(ch['text'] for _, chapters in parsed.sections for ch in chapters)
    assert 'Nee.' in bodies
