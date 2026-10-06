from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from quietwriter.docx_io import parse_docx_book
from quietwriter.i18n import set_locale


def _set_default_language(doc: Document, value: str):
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        doc.styles.element.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.append(rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), value)


def test_docx_import_keeps_only_primary_language_for_book_details(tmp_path: Path):
    path = tmp_path / "english-us.docx"
    doc = Document()
    _set_default_language(doc, "en-US")
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("Body")
    doc.save(path)
    assert parse_docx_book(path).language == "en"


def test_docx_import_warnings_follow_ui_locale(tmp_path: Path):
    path = tmp_path / "table.docx"
    doc = Document()
    doc.add_heading("Chapter", level=1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    doc.save(path)
    try:
        set_locale("en")
        warnings = parse_docx_book(path).warnings
        assert warnings == ("1 table(s) were not imported.",)
    finally:
        set_locale("nl")
