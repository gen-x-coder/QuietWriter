from __future__ import annotations

import base64
import json
from pathlib import Path

from docx import Document
from docx.shared import Mm

from quietwriter.docx_io import read_docx_import_document
from quietwriter.media.markup import find_image_references
from quietwriter.storage import Library


_PNG_1X1 = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZtJkAAAAASUVORK5CYII='
)


def _make_docx_with_image(tmp_path: Path, *, repeated: bool = False, inline_text: bool = False) -> Path:
    image_path = tmp_path / 'illustratie.png'
    image_path.write_bytes(_PNG_1X1)
    path = tmp_path / 'met-afbeelding.docx'

    doc = Document()
    doc.core_properties.title = 'Boek met beeld'
    doc.add_heading('Hoofdstuk Een', level=1)
    doc.add_paragraph('Voor de afbeelding.')
    p = doc.add_paragraph('Tekst met afbeelding: ' if inline_text else '')
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Mm(20))
    shape._inline.docPr.set('descr', 'Rode ballon')
    if repeated:
        p2 = doc.add_paragraph()
        p2.add_run().add_picture(str(image_path), width=Mm(20))
    doc.add_paragraph('Na de afbeelding.')
    doc.save(path)
    return path


def test_docx_reader_carries_supported_image_as_neutral_asset_and_block(tmp_path: Path):
    source = _make_docx_with_image(tmp_path)
    imported = read_docx_import_document(source)

    assert len(imported.assets) == 1
    asset = imported.assets[0]
    assert asset.media_type == 'image/png'
    assert asset.data == _PNG_1X1

    blocks = imported.sections[0].chapters[0].blocks
    image_blocks = [block for block in blocks if block.kind == 'image']
    assert len(image_blocks) == 1
    assert image_blocks[0].asset_key == asset.key
    assert image_blocks[0].alt == 'Rode ballon'
    assert not any('niet geïmporteerd' in warning and 'afbeeld' in warning.lower() for warning in imported.warnings)


def test_docx_book_import_publishes_image_in_media_store_and_chapter_reference(tmp_path: Path):
    source = _make_docx_with_image(tmp_path)
    lib = Library(tmp_path / 'workspace')

    book, warnings = lib.import_docx_book(source)
    assert warnings == ()

    manifest = json.loads((book.path / 'assets' / 'manifest.json').read_text(encoding='utf-8'))
    assert len(manifest['images']) == 1
    row = next(iter(manifest['images'].values()))
    asset_path = book.path / row['file']
    assert asset_path.read_bytes() == _PNG_1X1

    chapter = book.sections[0].chapters[0]
    source_text = (book.path / chapter.file).read_text(encoding='utf-8')
    refs = find_image_references(source_text)
    assert len(refs) == 1
    assert refs[0].alt == 'Rode ballon'
    assert refs[0].path.startswith('../assets/images/')


def test_repeated_same_docx_image_is_deduplicated_but_both_blocks_remain(tmp_path: Path):
    source = _make_docx_with_image(tmp_path, repeated=True)
    lib = Library(tmp_path / 'workspace')

    book, _warnings = lib.import_docx_book(source)
    manifest = json.loads((book.path / 'assets' / 'manifest.json').read_text(encoding='utf-8'))
    assert len(manifest['images']) == 1

    chapter = book.sections[0].chapters[0]
    source_text = (book.path / chapter.file).read_text(encoding='utf-8')
    refs = find_image_references(source_text)
    assert len(refs) == 2
    assert refs[0].path == refs[1].path


def test_inline_word_image_is_preserved_as_block_with_explicit_warning(tmp_path: Path):
    source = _make_docx_with_image(tmp_path, inline_text=True)
    imported = read_docx_import_document(source)

    blocks = imported.sections[0].chapters[0].blocks
    assert any(block.visible_text.startswith('Tekst met afbeelding:') for block in blocks)
    assert any(block.kind == 'image' for block in blocks)
    assert any('losse afbeeldingsblokken' in warning for warning in imported.warnings)


def test_image_write_failure_rolls_back_entire_docx_import(tmp_path: Path, monkeypatch):
    from quietwriter.media.store import MediaStore, MediaError

    source = _make_docx_with_image(tmp_path)
    lib = Library(tmp_path / 'workspace')

    def fail_import(*_args, **_kwargs):
        raise MediaError('gesimuleerde mediafout')

    monkeypatch.setattr(MediaStore, 'import_image_bytes', fail_import)

    import pytest
    with pytest.raises(MediaError):
        lib.import_docx_book(source)

    visible = [p for p in lib.books_dir.iterdir() if p.is_dir() and not p.name.startswith('.')]
    staging = [p for p in lib.books_dir.iterdir() if p.is_dir() and p.name.startswith('.import-')]
    assert visible == []
    assert staging == []
    assert lib.list_books() == []
