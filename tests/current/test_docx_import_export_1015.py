from __future__ import annotations

from io import BytesIO
from pathlib import Path
import struct
import zlib
import zipfile

from docx import Document
from docx.enum.text import WD_BREAK

from quietwriter.docx_io import export_docx, parse_docx_book
from quietwriter.exporting.models import ExportAsset, ExportChapter, ExportDocument, ExportSection


def _png(width=1200, height=1800):
    raw = b''.join(b'\x00' + (b'\x88\xaa\xcc' * width) for _ in range(height))
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw, 9)) + chunk(b'IEND', b'')


def _export_document():
    cover = ExportAsset(id='cover', filename='cover.png', media_type='image/png', data=_png(900, 1400), role='cover')
    huge = ExportAsset(id='inline-img', filename='huge.png', media_type='image/png', data=_png(1200, 1800), role='inline')
    return ExportDocument(
        id='b', title='Boek', author='Auteur', language='nl_NL', slug='boek', metadata={},
        sections=(ExportSection(id='root', title='Manuscript', chapters=(
            ExportChapter(id='c1', title='Eerste', markdown='Tekst\n![Groot](huge.png)'),
            ExportChapter(id='c2', title='Tweede', markdown='Meer tekst'),
        )),),
        front_matter=(), back_matter=(), assets=(cover, huge), cover_asset_id='cover',
    )


def test_docx_export_adds_cover_language_page_breaks_and_bounds_images(tmp_path):
    path = tmp_path / 'boek.docx'
    export_docx(_export_document(), path, {})
    doc = Document(path)
    assert len(doc.inline_shapes) == 2  # cover + manuscript image
    section = doc.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = section.page_height - section.top_margin - section.bottom_margin
    assert all(shape.width <= max_width and shape.height <= max_height for shape in doc.inline_shapes)

    chapter_paras = [p for p in doc.paragraphs if p.style.name == 'QuietWriter Chapter']
    assert [p.text for p in chapter_paras] == ['Eerste', 'Tweede']
    # Every chapter after the first starts on a new page without inserting a
    # separate break paragraph that could create an empty page after a large image.
    assert chapter_paras[1].paragraph_format.page_break_before is True

    with zipfile.ZipFile(path) as zf:
        styles = zf.read('word/styles.xml').decode('utf-8')
    assert 'w:lang' in styles and 'nl-NL' in styles


def test_layout_heavy_docx_uses_page_breaks_title_subtitle_and_large_display_text(tmp_path):
    path = tmp_path / 'layout.docx'
    doc = Document()
    doc.core_properties.title = 'Word Document'
    p = doc.add_paragraph('Whispers of the Soul')
    p.runs[0].font.size = __import__('docx').shared.Pt(30)
    doc.add_paragraph('A Collection of Modern Poems')
    doc.add_paragraph('JANE DOE')
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph('Copyright page')
    for _ in range(28):
        doc.add_paragraph('body')
    p = doc.add_paragraph('TABLE of Contents'); p.runs[0].font.size = __import__('docx').shared.Pt(24)
    doc.add_paragraph('Chapter 1 ........ 2')
    p = doc.add_paragraph('From the Author'); p.runs[0].font.size = __import__('docx').shared.Pt(24)
    p = doc.add_paragraph('DEDICATION'); p.runs[0].font.size = __import__('docx').shared.Pt(24)
    doc.add_paragraph('Dedicated text').add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph('CHAPTER ONE', style='Title')
    doc.add_paragraph('Reflections', style='Subtitle')
    doc.add_paragraph('Chapter body').add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading('THE DANCE OF TIME', level=1)
    doc.add_paragraph('Poem body')
    doc.save(path)

    parsed = parse_docx_book(path)
    assert parsed.title == 'Whispers of the Soul'
    titles = [c['title'] for _, chapters in parsed.sections for c in chapters]
    assert 'TABLE of Contents' in titles
    assert 'From the Author — DEDICATION' in titles
    assert 'CHAPTER ONE — Reflections' in titles
    assert 'THE DANCE OF TIME' in titles
