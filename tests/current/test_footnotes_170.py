"""1.7.0: voetnoten als ``^[tekst]`` op de plek in de tekst.

Niet getest: hoe Microsoft Word zelf het DOCX-bestand opent (hier alleen met
python-docx teruggelezen en handmatig met LibreOffice gecontroleerd), en hoe
e-readers de noten tonen (de EPUB wordt wel intern gevalideerd).
"""
from __future__ import annotations

import os
import tempfile
import zipfile
from pathlib import Path

import pytest

from quietwriter.document_view import (count_visible_words, inline_runs, note_groups, text_for_ai_context,
                                       text_for_language_tools, visible_text)
from quietwriter.manuscript_markup import escape_literal_text, escape_literal_typed_char, toggle_inline
from quietwriter.manuscript_syntax import parse_inline_spans

SOURCE = 'Het huis^[Gebouwd in **1902**.] stond aan de dijk.^[Tweede noot.]'


def _notes(text):
    return [(s.kind, text[s.content_start:s.content_end]) for s in parse_inline_spans(text) if s.kind == 'footnote']


def test_parser_reads_inline_notes_and_leaves_lookalikes_alone():
    assert _notes(SOURCE) == [('footnote', 'Gebouwd in **1902**.'), ('footnote', 'Tweede noot.')]
    assert _notes('a ^[nooit gesloten') == []
    assert _notes(r'a \^[geen noot]') == []
    assert _notes('`^[code]`') == []
    assert _notes(r'x^[een \] haak] y') == [('footnote', r'een \] haak')]
    assert _notes('^[a]^[b]') == [('footnote', 'a'), ('footnote', 'b')]
    # Gewone haken in proza blijven gewone tekst.
    assert parse_inline_spans('Hij zei [lachend] nee.') == []


def test_typed_and_pasted_text_can_never_open_or_close_a_note():
    assert escape_literal_typed_char('^', 'x') == r'\^'
    assert escape_literal_typed_char(']', 'x') == r'\]'
    pasted = escape_literal_text('zie^[dit] niet')
    assert _notes(pasted) == [] and visible_text(pasted) == 'zie^[dit] niet'


def test_notes_are_counted_searched_and_given_to_the_meelezer_separately():
    assert count_visible_words(SOURCE) == 11
    assert 'huis Gebouwd' in visible_text(SOURCE)          # niet aan elkaar geplakt
    assert 'huisGebouwd' not in visible_text(SOURCE)
    tools = text_for_language_tools(SOURCE)
    assert len(tools) == len(SOURCE) and '^[' not in tools  # offsets blijven gelijk voor spelling
    assert text_for_ai_context(SOURCE) == 'Het huis [Voetnoot: Gebouwd in 1902.] stond aan de dijk. [Voetnoot: Tweede noot.]'


def test_note_groups_keep_formatting_inside_a_note():
    groups = note_groups(inline_runs(SOURCE))
    kinds = [kind for kind, _ in groups]
    assert kinds == ['text', 'note', 'text', 'note']
    first_note = groups[1][1]
    assert [(r.text, sorted(r.styles)) for r in first_note] == [('Gebouwd in ', []), ('1902', ['bold']), ('.', [])]


def test_selection_becomes_a_note_and_back():
    text = 'Het huis stond aan de dijk.'
    new, start, end = toggle_inline(text, 9, 14, 'footnote')
    assert new == 'Het huis ^[stond] aan de dijk.' and new[start:end] == 'stond'
    assert toggle_inline(new, start, end, 'footnote')[0] == text


def test_epub_and_pdf_number_notes_per_chapter():
    from quietwriter.exporting.markup import markdown_to_xhtml
    from quietwriter.exporting.pdf_exporter import _markdown_to_pdf_html
    body = markdown_to_xhtml(SOURCE + '\n\nNog een^[Derde.] alinea.')
    assert body.count('epub:type="noteref"') == 3
    assert 'href="#fn-3"' in body and 'id="fn-3"' in body and 'href="#fnref-1"' in body
    assert '<strong>1902</strong>' in body.split('section class="footnotes"', 1)[1]
    assert 'Gebouwd' not in body.split('<section', 1)[0]     # noottekst staat niet in de lopende tekst
    pdf = _markdown_to_pdf_html(SOURCE, {}, 600)
    assert '<sup class="noteref">1</sup>' in pdf and 'class="footnote"' in pdf


def test_epub_with_notes_passes_validation():
    from quietwriter.exporting.epub_exporter import export_epub
    from quietwriter.exporting.epub_validation import validate_epub_archive
    from quietwriter.exporting.models import ExportChapter, ExportDocument, ExportSection
    from quietwriter.exporting.settings import default_export_settings
    document = ExportDocument(
        id='12345678-1234-1234-1234-123456789abc', title='Noten', author='Ada', language='nl', slug='noten',
        metadata={}, sections=(ExportSection('root', 'Manuscript', (
            ExportChapter('c1', 'Een', SOURCE), ExportChapter('c2', 'Twee', 'Weer^[Opnieuw één.] verder.'),
        )),), front_matter=(), back_matter=(), assets=(), cover_asset_id=None, epub_isbn='')
    settings = default_export_settings(); settings['epub']['include_cover'] = False
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / 'noten.epub'
        export_epub(document, path, settings)
        validate_epub_archive(path)
        with zipfile.ZipFile(path) as zf:
            chapter_two = next(zf.read(n).decode() for n in zf.namelist() if n.endswith('chapter-001-002.xhtml'))
            css = zf.read('EPUB/styles/book.css').decode()
    assert 'xmlns:epub=' in chapter_two and 'id="fn-1"' in chapter_two   # nummering per hoofdstuk
    assert 'section.footnotes' in css


def test_docx_export_writes_real_word_footnotes_and_import_reads_them_back(tmp_path):
    from docx import Document
    from quietwriter.docx_io import _append_markdown, _ensure_styles, paragraph_to_import_block
    from quietwriter.import_document import serialize_import_block

    class Export:
        assets = []

    doc = Document(); _ensure_styles(doc)
    _append_markdown(doc, SOURCE, Export())
    path = tmp_path / 'noten.docx'
    doc.save(str(path))
    with zipfile.ZipFile(path) as zf:
        footnotes = zf.read('word/footnotes.xml').decode()
        document_xml = zf.read('word/document.xml').decode()
        settings_xml = zf.read('word/settings.xml').decode()
        types = zf.read('[Content_Types].xml').decode()
    assert footnotes.count('<w:footnote ') >= 4      # twee scheidingsnoten + twee noten
    assert document_xml.count('w:footnoteReference') == 2
    assert 'footnotes+xml' in types and '<w:footnotePr>' in settings_xml
    assert settings_xml.index('<w:footnotePr>') < settings_xml.index('<w:compat')   # schemavolgorde
    reread = Document(str(path))
    assert serialize_import_block(paragraph_to_import_block(reread.paragraphs[0])) == SOURCE


def test_submission_export_also_uses_word_footnotes(tmp_path):
    from quietwriter.exporting.submission_exporter import export_submission
    from quietwriter.exporting.models import ExportChapter, ExportDocument, ExportSection
    document = ExportDocument(
        id='x', title='Noten', author='Ada Auteur', language='nl', slug='noten', metadata={},
        sections=(ExportSection('root', 'Manuscript', (ExportChapter('c1', 'Een', SOURCE),)),),
        front_matter=(), back_matter=(), assets=(), cover_asset_id=None, epub_isbn='')
    out = export_submission(document, tmp_path / 'm.docx')
    with zipfile.ZipFile(out) as zf:
        assert zf.read('word/document.xml').decode().count('w:footnoteReference') == 2


def _qt():
    pytest.importorskip('PySide6')
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.mark.qt
def test_ctrl_alt_f_puts_a_number_in_the_text_and_the_cursor_in_the_note_strip():
    # Feedback Lucas (dev9): "de cursor wordt klein maar verder zie ik niets". Nu zoals
    # in ONLYOFFICE: nummer in de tekst, onder het blad een streep met "1" en de cursor.
    app = _qt()
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from quietwriter.ui.footnote_strip import FootnoteStrip
    from quietwriter.ui.manuscript_editor import ManuscriptEditor
    editor = ManuscriptEditor()
    strip = FootnoteStrip(editor)
    editor.attach_footnote_strip(strip)
    editor.resize(900, 600); editor.show()
    editor.setPlainText('Het huis stond.')
    cursor = editor.textCursor(); cursor.setPosition(8); editor.setTextCursor(cursor)
    QTest.keyClick(editor, Qt.Key_F, Qt.ControlModifier | Qt.AltModifier)
    assert editor.toPlainText() == 'Het huis^[] stond.'
    assert not strip.isHidden() and len(strip.rows) == 1 and strip.rows[0].number.text() == '1'
    note = strip.rows[0].edit
    QTest.keyClicks(note, 'Uit 1902]')
    assert editor.toPlainText() == r'Het huis^[Uit 1902\]] stond.'      # ] in de noot sluit hem niet
    QTest.keyClick(note, Qt.Key_Return)                                  # terug naar de tekst, na de noot
    assert editor.textCursor().position() == editor.footnote_spans()[0][3]
    QTest.keyClicks(editor, '!')
    assert editor.toPlainText() == r'Het huis^[Uit 1902\]]! stond.'
    assert _notes(editor.toPlainText()) == [('footnote', r'Uit 1902\]')]
    # Een tweede noot ervoor wordt nummer 1; de strook nummert mee.
    cursor = editor.textCursor(); cursor.setPosition(3); editor.setTextCursor(cursor)
    QTest.keyClick(editor, Qt.Key_F, Qt.ControlModifier | Qt.AltModifier)
    assert [row.number.text() for row in strip.rows] == ['1', '2']
    assert strip.rows[1].edit.toPlainText() == r'Uit 1902\]'
    editor.close()
    assert app is not None


@pytest.mark.qt
def test_cursor_skips_the_collapsed_note_and_backspace_removes_it_whole():
    app = _qt()
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from quietwriter.ui.manuscript_editor import ManuscriptEditor
    editor = ManuscriptEditor()
    editor.setPlainText('Huis^[Noot.] staat.')
    cursor = editor.textCursor(); cursor.setPosition(4); editor.setTextCursor(cursor)
    QTest.keyClick(editor, Qt.Key_Right)           # één stap: over de hele noot heen
    assert editor.textCursor().position() == len('Huis^[Noot.]')
    QTest.keyClick(editor, Qt.Key_Left)
    assert editor.textCursor().position() == 4
    QTest.keyClick(editor, Qt.Key_Right)
    QTest.keyClick(editor, Qt.Key_Backspace)        # nummer weg = noot weg
    assert editor.toPlainText() == 'Huis staat.'
    editor.undo()
    assert editor.toPlainText() == 'Huis^[Noot.] staat.'
    # De noottekst is in de tekst ingeklapt: (bijna) geen breedte.
    editor.presentation_highlighter.rehighlight()
    block = editor.document().firstBlock()
    sizes = {r.start: r.format.fontPointSize() for r in block.layout().formats()}
    assert sizes.get(len('Huis^[')) == 1.0
    assert app is not None


@pytest.mark.qt
def test_editor_page_has_the_note_strip_under_the_page(tmp_path):
    app = _qt()
    from PySide6.QtCore import QSettings
    from quietwriter.storage import Library
    from quietwriter.ui.main_window import MainWindow
    lib = Library(tmp_path / 'ws')
    book = lib.create_book('Noten')
    settings = QSettings(str(tmp_path / 's.ini'), QSettings.IniFormat)
    settings.setValue('workspace', str(lib.root))
    window = MainWindow(settings, lib, [])
    window.resize(1100, 700); window.show(); window.open_book(book)
    page = window.editor_page
    page.editor.setPlainText('Een^[Eerste.] en twee^[Tweede.].')
    page.footnote_strip.refresh()
    app.processEvents()
    strip = page.footnote_strip
    assert strip.parent() is page.editor and not strip.isHidden()
    assert [row.edit.toPlainText() for row in strip.rows] == ['Eerste.', 'Tweede.']
    assert strip.geometry().top() >= page.editor.viewport().geometry().bottom()
    page.editor.setPlainText('Geen noten meer.')
    strip.refresh()
    assert strip.isHidden()
    window.close()
