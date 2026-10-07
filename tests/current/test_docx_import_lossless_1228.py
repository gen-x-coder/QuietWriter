from __future__ import annotations

import base64
import os
import time
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from quietwriter.docx_io import read_docx_import_document
from quietwriter.import_document import serialize_import_chapter
from quietwriter.storage import Library


_PNG_1X1 = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZtJkAAAAASUVORK5CYII='
)


def _all_source(imported) -> str:
    return '\n'.join(
        serialize_import_chapter(chapter)
        for section in imported.sections
        for chapter in section.chapters
    )


def test_soft_and_page_breaks_preserve_text_without_creating_hidden_paragraph_loss(tmp_path: Path):
    path = tmp_path / 'breaks.docx'
    doc = Document()
    doc.add_heading('Hoofdstuk', level=1)
    p = doc.add_paragraph()
    r = p.add_run('Voor')
    r.add_break()
    r.add_text('Na zacht')
    r.add_break(WD_BREAK.PAGE)
    r.add_text('Na pagina')
    doc.save(path)

    imported = read_docx_import_document(path)
    source = _all_source(imported)

    assert 'Voor\u2028Na zacht\u2028Na pagina' in source
    assert any('pagina-einde' in warning.lower() for warning in imported.warnings)


def test_table_rows_are_preserved_as_plain_readable_text(tmp_path: Path):
    path = tmp_path / 'table.docx'
    doc = Document()
    doc.add_heading('Hoofdstuk', level=1)
    doc.add_paragraph('Voor tabel')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Naam'
    table.cell(0, 1).text = 'Waarde'
    table.cell(1, 0).text = 'Rood'
    table.cell(1, 1).text = '42'
    doc.add_paragraph('Na tabel')
    doc.save(path)

    imported = read_docx_import_document(path)
    source = _all_source(imported)

    assert 'Naam · Waarde' in source
    assert 'Rood · 42' in source
    assert source.index('Voor tabel') < source.index('Naam · Waarde') < source.index('Na tabel')
    assert any('tabel' in warning.lower() and 'tekst' in warning.lower() for warning in imported.warnings)


def test_field_result_tracked_insertion_and_textbox_text_are_not_silently_lost(tmp_path: Path):
    path = tmp_path / 'ooxml.docx'
    doc = Document()
    doc.add_heading('Hoofdstuk', level=1)
    p = doc.add_paragraph('Basis ')

    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fld_run = OxmlElement('w:r')
    fld_text = OxmlElement('w:t')
    fld_text.text = '7'
    fld_run.append(fld_text)
    fld.append(fld_run)
    p._p.append(fld)

    ins = OxmlElement('w:ins')
    ins_run = OxmlElement('w:r')
    ins_text = OxmlElement('w:t')
    ins_text.text = ' ingevoegd'
    ins_run.append(ins_text)
    ins.append(ins_run)
    p._p.append(ins)

    txbx = OxmlElement('w:txbxContent')
    tx_p = OxmlElement('w:p')
    tx_r = OxmlElement('w:r')
    tx_t = OxmlElement('w:t')
    tx_t.text = ' tekstvakinhoud'
    tx_r.append(tx_t)
    tx_p.append(tx_r)
    txbx.append(tx_p)
    p._p.append(txbx)
    doc.save(path)

    imported = read_docx_import_document(path)
    source = _all_source(imported)

    assert 'Basis 7 ingevoegd tekstvakinhoud' in source
    assert any('tekstvakken' in warning.lower() for warning in imported.warnings)
    assert any('velden' in warning.lower() for warning in imported.warnings)
    assert any('wijzigingen' in warning.lower() for warning in imported.warnings)


def test_corrupt_embedded_image_is_skipped_with_warning_instead_of_aborting_import(tmp_path: Path):
    image = tmp_path / 'img.png'
    image.write_bytes(_PNG_1X1)
    original = tmp_path / 'original.docx'
    corrupt = tmp_path / 'corrupt.docx'

    doc = Document()
    doc.add_heading('Hoofdstuk', level=1)
    doc.add_paragraph('Tekst blijft bestaan.')
    doc.add_paragraph().add_run().add_picture(str(image))
    doc.add_paragraph('Ook deze tekst blijft.')
    doc.save(original)

    with zipfile.ZipFile(original, 'r') as src, zipfile.ZipFile(corrupt, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith('word/media/'):
                data = b'geen geldige afbeelding'
            dst.writestr(info, data)

    lib = Library(tmp_path / 'workspace')
    book, warnings = lib.import_docx_book(corrupt)
    chapter_text = '\n'.join(
        (book.path / chapter.file).read_text(encoding='utf-8')
        for section in book.sections for chapter in section.chapters
    )

    assert 'Tekst blijft bestaan.' in chapter_text
    assert 'Ook deze tekst blijft.' in chapter_text
    assert any('beschadigde' in warning.lower() and 'afbeelding' in warning.lower() for warning in warnings)
    manifest = (book.path / 'assets' / 'manifest.json').read_text(encoding='utf-8')
    assert '"images": {}' in manifest


def test_library_startup_removes_only_old_abandoned_import_staging(tmp_path: Path):
    root = tmp_path / 'workspace'
    books = root / 'books'
    books.mkdir(parents=True)
    old_stage = books / '.import-old'
    fresh_stage = books / '.import-fresh'
    old_stage.mkdir()
    fresh_stage.mkdir()
    old_time = time.time() - 2 * 24 * 60 * 60
    os.utime(old_stage, (old_time, old_time))

    Library(root)

    assert not old_stage.exists()
    assert fresh_stage.exists()
