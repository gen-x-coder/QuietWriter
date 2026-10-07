from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from quietwriter.docx_io import parse_docx_book
from quietwriter.i18n import set_locale


def _set_default_language(doc: Document, value: str):
    styles = doc.styles.element
    defaults = styles.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults')
        styles.insert(0, defaults)
    rpr_default = defaults.find(qn('w:rPrDefault'))
    if rpr_default is None:
        rpr_default = OxmlElement('w:rPrDefault')
        defaults.append(rpr_default)
    rpr = rpr_default.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        rpr_default.append(rpr)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), value)


def _long_page_document(path: Path):
    doc = Document()
    doc.add_paragraph('Mijn Roman', style='Title').add_run().add_break(WD_BREAK.PAGE)
    for idx in range(3):
        text = (f'Pagina {idx + 1} begint met een lange alinea die beslist geen kop is. ' * 4).strip()
        doc.add_paragraph(text).add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)


def test_page_fallback_titles_are_book_wide_and_translated(tmp_path: Path):
    path = tmp_path / 'fallbacks.docx'
    _long_page_document(path)

    set_locale('nl')
    parsed_nl = parse_docx_book(path)
    nl_titles = [ch['title'] for _, chapters in parsed_nl.sections for ch in chapters]
    assert nl_titles[-3:] == ['Deel 2', 'Deel 3', 'Deel 4']

    set_locale('en')
    parsed_en = parse_docx_book(path)
    en_titles = [ch['title'] for _, chapters in parsed_en.sections for ch in chapters]
    assert en_titles[-3:] == ['Part 2', 'Part 3', 'Part 4']
    set_locale('nl')


def test_docx_import_reads_and_normalizes_default_document_language(tmp_path: Path):
    path = tmp_path / 'english.docx'
    doc = Document()
    _set_default_language(doc, 'en_us')
    doc.add_heading('Chapter One', level=1)
    doc.add_paragraph('English body text.')
    doc.save(path)

    parsed = parse_docx_book(path)
    assert parsed.language == 'en'
