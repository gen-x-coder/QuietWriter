from pathlib import Path
import os
import pytest

pytest.importorskip('PySide6')
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from quietwriter.storage import Library
from quietwriter.ui.book_details import BookDetailsPage


def _set_default_language(doc: Document, value: str):
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults')
        doc.styles.element.insert(0, defaults)
    rpr_default = defaults.find(qn('w:rPrDefault'))
    if rpr_default is None:
        rpr_default = OxmlElement('w:rPrDefault')
        defaults.append(rpr_default)
    rpr = rpr_default.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        rpr_default.append(rpr)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), value)


def test_import_en_us_stays_english_after_book_details_save(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    source = tmp_path / 'english-us.docx'
    doc = Document()
    _set_default_language(doc, 'en-US')
    doc.core_properties.title = 'English book'
    doc.add_heading('Chapter One', level=1)
    doc.add_paragraph('English body text.')
    doc.save(source)

    lib = Library(tmp_path / 'workspace')
    book, _warnings = lib.import_docx_book(source)
    assert book.metadata.get('language') == 'en'

    settings = QSettings(str(tmp_path / 'settings.ini'), QSettings.IniFormat)
    page = BookDetailsPage(lib, settings, book)
    assert page.language.currentData() == 'en'
    page.title_edit.setText('English book revised')
    assert page.save() is True

    reloaded = lib.load_book(book.path)
    assert reloaded.metadata.get('language') == 'en'
