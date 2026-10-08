from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from quietwriter.library_shelves import DEFAULT_SHELF_NAME, ShelfLibraryStore
from quietwriter.qwbook_io import export_qwbook, import_qwbook
from quietwriter.storage import Library


def test_new_library_creates_one_public_default_shelf_and_assigns_new_book(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Mijn boek')

    state = library.shelves.state()
    assert len(state.shelves) == 1
    assert state.shelves[0].name == DEFAULT_SHELF_NAME
    assert state.shelves[0].private is False
    assert state.default_shelf_id == state.shelves[0].id
    assert state.book_shelves[book.id] == state.default_shelf_id


def test_rename_uses_id_and_does_not_change_book_mapping(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    before = library.shelves.state()
    shelf_id = before.default_shelf_id

    after = library.shelves.rename_shelf(shelf_id, 'Lopende projecten')

    assert after.default_shelf_id == shelf_id
    assert after.shelf(shelf_id).name == 'Lopende projecten'
    assert after.book_shelves[book.id] == shelf_id


def test_default_shelf_cannot_be_private_or_deleted(tmp_path: Path):
    store = ShelfLibraryStore(tmp_path)
    state = store.state()

    with pytest.raises(ValueError):
        store.set_private(state.default_shelf_id, True)
    with pytest.raises(ValueError):
        store.delete_shelf(state.default_shelf_id)


def test_move_book_is_one_mapping_and_delete_shelf_returns_it_to_default(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    private = library.shelves.create_shelf('Onder pseudoniem', private=True)

    moved = library.shelves.move_book(book.id, private.id)
    assert moved.book_shelves[book.id] == private.id

    deleted = library.shelves.delete_shelf(private.id)
    assert deleted.book_shelves[book.id] == deleted.default_shelf_id
    assert deleted.shelf(private.id) is None


def test_trash_keeps_private_mapping_and_restore_preserves_it(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Geheim')
    private = library.shelves.create_shelf('Privé', private=True)
    library.shelves.move_book(book.id, private.id)

    library.track_book(book)
    library.delete_book(book)
    row = library.list_trashed_books()[0]
    assert library.shelves.state().book_shelves[book.id] == private.id

    restored = library.restore_trashed_book(row['path'])
    assert restored.id == book.id
    assert library.shelves.state().book_shelves[book.id] == private.id


def test_permanent_delete_forgets_mapping(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Weg')
    library.track_book(book)
    library.delete_book(book)
    row = library.list_trashed_books()[0]

    library.permanently_delete_trashed_book(row['path'])

    assert book.id not in library.shelves.state().book_shelves


def test_migration_maps_existing_live_and_trashed_books_without_touching_book_json(tmp_path: Path):
    library = Library(tmp_path)
    first = library.create_book('Eén')
    second = library.create_book('Twee')
    library.track_book(second)
    library.delete_book(second)

    # Simulate a 1.2 workspace: no library metadata yet.
    library.shelves.path.unlink()
    library.shelves.last_good_path.unlink(missing_ok=True)
    first_bytes = first.manifest_path.read_bytes()
    trashed_manifest = library.list_trashed_books()[0]['path'] / 'book.json'
    second_bytes = trashed_manifest.read_bytes()

    store = ShelfLibraryStore(tmp_path)
    state = store.state()

    assert state.book_shelves[first.id] == state.default_shelf_id
    assert state.book_shelves[second.id] == state.default_shelf_id
    assert first.manifest_path.read_bytes() == first_bytes
    assert trashed_manifest.read_bytes() == second_bytes


def test_qwbook_export_contains_no_library_metadata_and_import_assigns_local_default(tmp_path: Path):
    source_library = Library(tmp_path / 'source')
    book = source_library.create_book('Export')
    private = source_library.shelves.create_shelf('Privé', private=True)
    source_library.shelves.move_book(book.id, private.id)
    package = tmp_path / 'book.qwbook'

    export_qwbook(source_library, book, package)
    with zipfile.ZipFile(package, 'r') as archive:
        assert 'library.json' not in archive.namelist()
        for name in archive.namelist():
            if name.endswith('book.json'):
                text = archive.read(name).decode('utf-8')
                assert 'shelf_id' not in text
                assert 'book_shelves' not in text

    target_library = Library(tmp_path / 'target')
    imported = import_qwbook(target_library, package)
    target_state = target_library.shelves.state()
    assert target_state.book_shelves[imported.id] == target_state.default_shelf_id


def test_two_store_instances_replay_operations_on_latest_disk_state(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    store_a = ShelfLibraryStore(tmp_path)
    store_b = ShelfLibraryStore(tmp_path)
    private = store_a.create_shelf('Privé', private=True)

    # B may have been instantiated earlier, but every operation reads fresh disk state.
    store_b.rename_shelf(private.id, 'Pseudoniem')
    store_a.move_book(book.id, private.id)

    state = ShelfLibraryStore(tmp_path).state()
    assert state.shelf(private.id).name == 'Pseudoniem'
    assert state.shelf(private.id).private is True
    assert state.book_shelves[book.id] == private.id


def test_missing_mapping_is_default_in_normal_mode_but_hidden_in_demo(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    state = library.shelves.state()
    raw = json.loads(library.shelves.path.read_text(encoding='utf-8'))
    raw['book_shelves'].pop(book.id)
    library.shelves.path.write_text(json.dumps(raw), encoding='utf-8')

    assert library.shelves.effective_shelf_id(book.id, demo_mode=False) == state.default_shelf_id
    assert library.shelves.effective_shelf_id(book.id, demo_mode=True) is None
    # Reading does not repair/write the missing mapping.
    assert book.id not in json.loads(library.shelves.path.read_text(encoding='utf-8'))['book_shelves']


def test_private_book_is_hidden_in_demo_but_visible_normally(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Privé boek')
    private = library.shelves.create_shelf('Privé', private=True)
    library.shelves.move_book(book.id, private.id)

    assert book.id in library.shelves.visible_book_ids([book.id], demo_mode=False)
    assert book.id not in library.shelves.visible_book_ids([book.id], demo_mode=True)
    assert private.id not in library.shelves.visible_shelf_ids(demo_mode=True)


def test_corrupt_canonical_uses_last_good_normally_and_fails_closed_in_demo(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    # create_book produced both canonical and last-good.
    library.shelves.path.write_text('{broken', encoding='utf-8')

    loaded = library.shelves.load()
    assert loaded.source == 'last_good'
    assert book.id in library.shelves.visible_book_ids([book.id], demo_mode=False)
    assert library.shelves.visible_book_ids([book.id], demo_mode=True) == set()
    assert library.shelves.visible_shelf_ids(demo_mode=True) == ()


def test_corrupt_without_last_good_never_overwrites_source(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    library.shelves.last_good_path.unlink(missing_ok=True)
    corrupt = b'{not json at all'
    library.shelves.path.write_bytes(corrupt)

    loaded = library.shelves.load()

    assert loaded.source == 'emergency'
    assert library.shelves.path.read_bytes() == corrupt
    assert book.id in library.shelves.visible_book_ids([book.id], demo_mode=False)
    assert library.shelves.visible_book_ids([book.id], demo_mode=True) == set()


def test_missing_canonical_with_last_good_does_not_remigrate_or_leak_private_book(tmp_path: Path):
    library = Library(tmp_path)
    public = library.create_book('Publiek')
    private_book = library.create_book('Geheim')
    private = library.shelves.create_shelf('Privé', private=True)
    library.shelves.move_book(private_book.id, private.id)
    last_good_before = library.shelves.last_good_path.read_bytes()

    library.shelves.path.unlink()
    loaded = library.shelves.load()

    assert loaded.source == 'last_good'
    assert library.shelves.path.exists() is False
    assert loaded.state.shelf(private.id).private is True
    assert library.shelves.last_good_path.read_bytes() == last_good_before
    assert library.shelves.visible_book_ids([public.id, private_book.id], demo_mode=True) == set()


def test_default_shelf_id_is_deterministic_for_independent_migrations(tmp_path: Path):
    import shutil

    first_root = tmp_path / 'first'
    first = Library(first_root)
    books = [first.create_book(f'Boek {i}') for i in range(3)]
    first.shelves.path.unlink()
    first.shelves.last_good_path.unlink(missing_ok=True)

    second_root = tmp_path / 'second'
    shutil.copytree(first_root, second_root)

    first_store = ShelfLibraryStore(first_root)
    second_store = ShelfLibraryStore(second_root)
    first_loaded = first_store.load()
    second_loaded = second_store.load()

    assert first_loaded.source == second_loaded.source == 'pending_migration'
    assert not (first_root / 'library.json').exists()
    assert not (second_root / 'library.json').exists()
    assert first_loaded.state.default_shelf_id == second_loaded.state.default_shelf_id
    assert first_store._render(first_loaded.state) == second_store._render(second_loaded.state)
    assert {b.id for b in books} <= set(first_loaded.state.book_shelves)


def test_create_book_remains_successful_when_shelf_metadata_is_corrupt(tmp_path: Path):
    library = Library(tmp_path)
    library.create_book('Bestaand')
    library.shelves.path.write_text('{kapot', encoding='utf-8')

    created = library.create_book('Nieuw')

    assert created.title == 'Nieuw'
    assert any(book.id == created.id for book in library.list_books())
    warnings = library.pop_shelf_metadata_warnings()
    assert len(warnings) == 1
    assert 'plankindeling' in warnings[0]


def test_create_book_remains_successful_for_future_shelf_schema(tmp_path: Path):
    library = Library(tmp_path)
    library.create_book('Bestaand')
    library.shelves.path.write_text(json.dumps({
        'schema_version': 2,
        'default_shelf_id': 'future',
        'shelves': [],
        'book_shelves': {},
    }), encoding='utf-8')

    created = library.create_book('Nieuw')

    assert any(book.id == created.id for book in library.list_books())
    assert library.pop_shelf_metadata_warnings()
    assert library.shelves.visible_book_ids([created.id], demo_mode=True) == set()


def test_future_schema_queries_fail_closed_without_raising(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    library.shelves.path.write_text(json.dumps({
        'schema_version': 2,
        'default_shelf_id': 'future',
        'shelves': [],
        'book_shelves': {},
    }), encoding='utf-8')

    loaded = library.shelves.load()
    assert loaded.source == 'future'
    assert library.shelves.visible_book_ids([book.id], demo_mode=True, loaded=loaded) == set()
    assert library.shelves.visible_shelf_ids(demo_mode=True, loaded=loaded) == ()
    assert library.shelves.effective_shelf_id(book.id, demo_mode=True, loaded=loaded) is None


def test_unknown_shelf_mapping_is_tolerated_for_queries_but_blocks_writes(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Boek')
    raw = json.loads(library.shelves.path.read_text(encoding='utf-8'))
    raw['book_shelves'][book.id] = 'missing-shelf'
    library.shelves.path.write_text(json.dumps(raw), encoding='utf-8')

    loaded = library.shelves.load()
    assert loaded.source == 'canonical'
    assert library.shelves.effective_shelf_id(book.id, demo_mode=False, loaded=loaded) == loaded.state.default_shelf_id
    assert library.shelves.effective_shelf_id(book.id, demo_mode=True, loaded=loaded) is None
    with pytest.raises(Exception):
        library.shelves.create_shelf('Kan niet schrijven')


def test_explicit_restore_from_last_good_recreates_canonical(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Geheim')
    private = library.shelves.create_shelf('Privé', private=True)
    library.shelves.move_book(book.id, private.id)
    expected = library.shelves.last_good_path.read_bytes()
    library.shelves.path.unlink()

    restored = library.shelves.restore_from_last_good()

    assert library.shelves.path.read_bytes() == expected
    assert restored.book_shelves[book.id] == private.id


def test_docx_import_succeeds_when_shelf_metadata_is_corrupt(tmp_path: Path):
    from docx import Document

    library = Library(tmp_path)
    library.create_book('Bestaand')
    library.shelves.path.write_text('{kapot', encoding='utf-8')
    source = tmp_path / 'import.docx'
    doc = Document()
    doc.add_heading('Importboek', 1)
    doc.add_paragraph('tekst')
    doc.save(source)

    imported, warnings = library.import_docx_book(source)

    assert imported in library.list_books()
    assert any('plankindeling' in warning for warning in warnings)


def test_qwbook_import_succeeds_when_shelf_metadata_is_corrupt(tmp_path: Path):
    source_library = Library(tmp_path / 'source')
    source_book = source_library.create_book('Reiziger')
    package = tmp_path / 'reiziger.qwbook'
    export_qwbook(source_library, source_book, package)

    target = Library(tmp_path / 'target')
    target.create_book('Bestaand')
    target.shelves.path.write_text('{kapot', encoding='utf-8')

    imported = import_qwbook(target, package)

    assert imported.id == source_book.id
    assert any(book.id == imported.id for book in target.list_books())
    assert target.pop_shelf_metadata_warnings()


def _shelf_process_worker(root: str, prefix: str, count: int):
    store = ShelfLibraryStore(Path(root))
    for index in range(count):
        store.create_shelf(f'{prefix}{index}')


def test_concurrent_processes_do_not_silently_lose_shelf_operations(tmp_path: Path):
    import multiprocessing

    Library(tmp_path).shelves.state()
    ctx = multiprocessing.get_context('spawn')
    processes = [
        ctx.Process(target=_shelf_process_worker, args=(str(tmp_path), prefix, 20))
        for prefix in ('X', 'Y')
    ]
    for process in processes:
        process.start()
    for process in processes:
        process.join(20)
        assert process.exitcode == 0

    names = {s.name for s in ShelfLibraryStore(tmp_path).state().shelves}
    assert {f'X{i}' for i in range(20)} <= names
    assert {f'Y{i}' for i in range(20)} <= names


def test_missing_metadata_queries_stay_pending_without_writing_and_demo_fails_closed(tmp_path: Path):
    library = Library(tmp_path)
    public = library.create_book('Publiek')
    private_book = library.create_book('Geheim')
    private = library.shelves.create_shelf('Privé', private=True)
    library.shelves.move_book(private_book.id, private.id)

    # Simulate a second computer that received the books but not workspace metadata.
    library.shelves.path.unlink()
    library.shelves.last_good_path.unlink(missing_ok=True)
    loaded = library.shelves.load()

    assert loaded.source == 'pending_migration'
    assert library.shelves.path.exists() is False
    assert public.id in library.shelves.visible_book_ids([public.id, private_book.id], loaded=loaded)
    assert library.shelves.visible_book_ids([public.id, private_book.id], demo_mode=True, loaded=loaded) == set()
    assert library.shelves.path.exists() is False


def test_first_explicit_shelf_operation_persists_pending_migration(tmp_path: Path):
    library = Library(tmp_path)
    book = library.create_book('Bestaand')
    library.shelves.path.unlink()
    library.shelves.last_good_path.unlink(missing_ok=True)

    assert library.shelves.load().source == 'pending_migration'
    assert not library.shelves.path.exists()

    shelf = library.shelves.create_shelf('Romans')

    assert library.shelves.path.exists()
    state = library.shelves.state()
    assert state.shelf(shelf.id).name == 'Romans'
    assert state.book_shelves[book.id] == state.default_shelf_id


def test_recovery_refuses_to_overwrite_future_schema(tmp_path: Path):
    library = Library(tmp_path)
    library.create_book('Boek')
    future = {
        'schema_version': 2,
        'default_shelf_id': 'future',
        'shelves': [],
        'book_shelves': {},
        'future_field': 'keep-me',
    }
    library.shelves.path.write_text(json.dumps(future), encoding='utf-8')

    with pytest.raises(Exception):
        library.shelves.restore_from_last_good()
    assert 'keep-me' in library.shelves.path.read_text(encoding='utf-8')

    with pytest.raises(Exception):
        library.shelves.reset_to_default()
    assert 'keep-me' in library.shelves.path.read_text(encoding='utf-8')


def test_lockfile_is_local_and_outside_workspace(tmp_path: Path):
    store = ShelfLibraryStore(tmp_path)
    assert tmp_path not in store.lock_path.parents
    store.create_shelf('Test')
    assert not store.lock_path.exists()
