from pathlib import Path

from docx import Document

from quietwriter.docx_io import read_docx_import_document, parse_docx_book
from quietwriter.import_document import (
    ImportBlock, ImportChapter, ImportDocument, ImportInlineRun, ImportSection,
    serialize_import_document,
)


def test_neutral_import_model_contains_semantics_not_quietwriter_markup():
    imported = ImportDocument(
        title='Boek',
        sections=(ImportSection(None, (
            ImportChapter('Hoofdstuk', (
                ImportBlock('paragraph', (
                    ImportInlineRun('Letterlijk **sterretjes** '),
                    ImportInlineRun('vet', frozenset({'bold'})),
                )),
            )),
        )),),
    )
    block = imported.sections[0].chapters[0].blocks[0]
    assert block.visible_text == 'Letterlijk **sterretjes** vet'
    assert block.runs[0].styles == frozenset()
    assert block.runs[1].styles == frozenset({'bold'})

    serialized = serialize_import_document(imported)
    source = serialized[0][1][0]['text']
    assert r'\*\*sterretjes\*\*' in source
    assert '**vet**' in source


def test_docx_reader_returns_neutral_blocks_before_source_serialization(tmp_path: Path):
    path = tmp_path / 'neutral.docx'
    doc = Document()
    doc.add_heading('Eerste hoofdstuk', level=1)
    p = doc.add_paragraph('Gewoon *316 en ')
    run = p.add_run('vet')
    run.bold = True
    doc.save(path)

    imported = read_docx_import_document(path)
    chapter = imported.sections[0].chapters[0]
    assert chapter.title == 'Eerste hoofdstuk'
    assert chapter.blocks[0].kind == 'paragraph'
    assert chapter.blocks[0].visible_text == 'Gewoon *316 en vet'
    assert chapter.blocks[0].runs[-1].styles == frozenset({'bold'})

    # The legacy/storage boundary still receives safe QuietWriter source.
    parsed = parse_docx_book(path)
    source = parsed.sections[0][1][0]['text']
    assert r'\*316' in source
    assert '**vet**' in source
