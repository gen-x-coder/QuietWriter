from pathlib import Path
import os
import pytest

pytest.importorskip('PySide6')
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from docx import Document

from quietwriter.storage import Library
from quietwriter.ui.main_window import MainWindow


def _window(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    lib = Library(tmp_path / 'workspace')
    settings = QSettings(str(tmp_path / 'settings.ini'), QSettings.IniFormat)
    settings.setValue('spell_enabled', False)
    settings.setValue('ai_enabled', False)
    settings.setValue('advanced_options', False)
    window = MainWindow(settings, lib, [])
    return app, lib, window


def test_export_page_exposes_docx_and_can_create_real_docx(tmp_path: Path, monkeypatch):
    _app, lib, window = _window(tmp_path)
    try:
        book = lib.create_book('Word boek')
        lib.save_chapter(book, book.sections[0].chapters[0], 'Tekst met **vet**.')
        window.open_book(book)
        page = window.export_page
        page.set_book(window._active_book)
        page.docx_button.setChecked(True)
        page._format_changed()
        assert page.format_name == 'docx'
        assert 'DOCX' in page.export_button.text()
        out = tmp_path / 'out'
        out.mkdir()
        window.settings.setValue('export/output_dir', str(out))
        monkeypatch.setattr(QMessageBox, 'information', lambda *a, **k: QMessageBox.Ok)
        page._export()
        files = list(out.glob('*.docx'))
        assert len(files) == 1
        reopened = Document(str(files[0]))
        assert 'Word boek' in '\n'.join(p.text for p in reopened.paragraphs)
    finally:
        window.close()


def test_main_window_import_dispatches_docx_and_opens_book(tmp_path: Path, monkeypatch):
    _app, lib, window = _window(tmp_path)
    source = tmp_path / 'import.docx'
    doc = Document(); doc.core_properties.title = 'Geïmporteerd'; doc.add_heading('Hoofdstuk 1', level=1); doc.add_paragraph('Tekst'); doc.save(source)
    try:
        monkeypatch.setattr(QFileDialog, 'getOpenFileName', lambda *a, **k: (str(source), 'Word document (*.docx)'))
        window.import_book()
        assert window._active_book is not None
        assert window._active_book.title == 'Geïmporteerd'
        assert lib.read_chapter(window._active_book, window._active_book.sections[0].chapters[0]) == 'Tekst'
    finally:
        window.close()
