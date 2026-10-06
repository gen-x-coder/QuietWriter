from pathlib import Path

from docx import Document

from quietwriter.docx_io import export_docx, parse_docx_book
from quietwriter.exporting.models import ExportChapter, ExportDocument, ExportSection


def _document():
    return ExportDocument(
        id='book-1', title='DOCX Boek', author='Auteur', language='nl', slug='docx-boek', metadata={},
        sections=(
            ExportSection(id='part-1', title='Deel I', chapters=(
                ExportChapter(id='c1', title='Hoofdstuk 1', markdown='Gewoon **vet** en *schuin*.\n\n***\n<u>onder</u> en ~~weg~~ en `code`.'),
                ExportChapter(id='c2', title='Hoofdstuk 2', markdown='- één\n- twee\n> citaat'),
            )),
        ),
        front_matter=(), back_matter=(),
    )


def test_docx_export_and_reimport_preserves_structure_and_supported_inline_markup(tmp_path):
    destination = tmp_path / 'boek.docx'
    export_docx(_document(), destination, {})
    assert destination.exists()

    parsed = parse_docx_book(destination)
    assert parsed.title == 'DOCX Boek'
    assert parsed.author == 'Auteur'
    assert len(parsed.sections) == 1
    section_title, chapters = parsed.sections[0]
    assert section_title == 'Deel I'
    assert [c['title'] for c in chapters] == ['Hoofdstuk 1', 'Hoofdstuk 2']
    assert chapters[0]['text'] == 'Gewoon **vet** en *schuin*.\n\n***\n<u>onder</u> en ~~weg~~ en `code`.'
    assert chapters[1]['text'] == '- één\n- twee\n> citaat'


def test_external_docx_heading1_becomes_chapters_when_no_heading2(tmp_path):
    path = tmp_path / 'roman.docx'
    doc = Document()
    doc.core_properties.title = 'Roman'
    doc.add_heading('Eerste', level=1)
    p = doc.add_paragraph(); p.add_run('Dit is '); r = p.add_run('vet'); r.bold = True; p.add_run('.')
    doc.add_heading('Tweede', level=1)
    doc.add_paragraph('Volgende tekst')
    doc.save(path)

    parsed = parse_docx_book(path)
    assert parsed.sections[0][0] is None
    assert [c['title'] for c in parsed.sections[0][1]] == ['Eerste', 'Tweede']
    assert parsed.sections[0][1][0]['text'] == 'Dit is **vet**.'


def test_external_docx_heading1_and_heading2_map_to_sections_and_chapters(tmp_path):
    path = tmp_path / 'delen.docx'
    doc = Document()
    doc.add_heading('Deel een', level=1)
    doc.add_heading('Hoofdstuk A', level=2)
    doc.add_paragraph('A')
    doc.add_heading('Hoofdstuk B', level=2)
    doc.add_paragraph('B')
    doc.add_heading('Deel twee', level=1)
    doc.add_heading('Hoofdstuk C', level=2)
    doc.add_paragraph('C')
    doc.save(path)

    parsed = parse_docx_book(path)
    assert [section for section, _ in parsed.sections] == ['Deel een', 'Deel twee']
    assert [c['title'] for c in parsed.sections[0][1]] == ['Hoofdstuk A', 'Hoofdstuk B']
    assert [c['title'] for c in parsed.sections[1][1]] == ['Hoofdstuk C']


def test_docx_import_reports_tables_and_images_instead_of_silently_claiming_support(tmp_path):
    path = tmp_path / 'table.docx'
    doc = Document()
    doc.add_heading('Hoofdstuk', level=1)
    doc.add_paragraph('Tekst')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Niet ondersteund'
    doc.save(path)
    parsed = parse_docx_book(path)
    assert any('tabel' in warning for warning in parsed.warnings)
