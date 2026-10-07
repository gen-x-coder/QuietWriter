from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import hashlib
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm
from docx.text.run import Run
from docx.table import Table

from .document_view import content_inline_runs, inline_runs, parse_document
from .manuscript_markup import serialize_block_source
from .import_document import (
    ImportAsset, ImportBlock, ImportChapter, ImportDocument, ImportInlineRun, ImportSection,
    serialize_import_block, serialize_import_document,
)
from .i18n import tr


@dataclass(frozen=True)
class DocxImportResult:
    title: str
    sections: tuple[tuple[str | None, tuple[dict, ...]], ...]
    author: str = ''
    language: str = ''
    warnings: tuple[str, ...] = ()



def _normalize_language_tag(value: str | None) -> str:
    raw = str(value or '').strip().replace('_', '-')
    if not raw:
        return ''
    parts = [part for part in raw.split('-') if part]
    if not parts:
        return ''
    if len(parts) == 1:
        return parts[0].lower()
    return '-'.join([parts[0].lower()] + [part.upper() if len(part) in (2, 3) else part for part in parts[1:]])


def _document_default_language(doc: Document) -> str:
    """Return Word's default document language, if explicitly stored."""
    try:
        defaults = doc.styles.element.find(qn('w:docDefaults'))
        if defaults is None:
            return ''
        rpr_default = defaults.find(qn('w:rPrDefault'))
        if rpr_default is None:
            return ''
        rpr = rpr_default.find(qn('w:rPr'))
        if rpr is None:
            return ''
        lang = rpr.find(qn('w:lang'))
        if lang is None:
            return ''
        value = lang.get(qn('w:val')) or lang.get(qn('w:eastAsia')) or ''
        normalized = _normalize_language_tag(value)
        return normalized.split('-', 1)[0] if normalized else ''
    except Exception:
        return ''

def _style_name(paragraph) -> str:
    try:
        return str(paragraph.style.name or '')
    except Exception:
        return ''


def _style_chain(style):
    """Yield a Word style and its basedOn chain without looping forever."""
    seen: set[int] = set()
    current = style
    while current is not None and id(current) not in seen:
        seen.add(id(current))
        yield current
        try:
            current = current.base_style
        except Exception:
            current = None


def _style_font_value(style, attr: str):
    for item in _style_chain(style):
        try:
            value = getattr(item.font, attr)
        except Exception:
            value = None
        if value is not None:
            return value
    return None


def _effective_run_bool(run, paragraph, direct_attr: str, font_attr: str | None = None, *, direct_on_font: bool = False) -> bool:
    """Resolve direct, character-style and paragraph-style Word formatting."""
    attr = font_attr or direct_attr
    try:
        direct = getattr(run.font if direct_on_font else run, direct_attr)
    except Exception:
        direct = None
    if direct is not None:
        return bool(direct)
    try:
        value = _style_font_value(run.style, attr)
    except Exception:
        value = None
    if value is not None:
        return bool(value)
    try:
        value = _style_font_value(paragraph.style, attr)
    except Exception:
        value = None
    return bool(value) if value is not None else False


def _effective_font_name(run, paragraph) -> str:
    try:
        direct = run.font.name
    except Exception:
        direct = None
    if direct:
        return str(direct).lower()
    try:
        value = _style_font_value(run.style, 'name')
    except Exception:
        value = None
    if value:
        return str(value).lower()
    try:
        value = _style_font_value(paragraph.style, 'name')
    except Exception:
        value = None
    return str(value or '').lower()


def _paragraph_style_names(paragraph) -> tuple[str, ...]:
    try:
        return tuple(str(style.name or '') for style in _style_chain(paragraph.style))
    except Exception:
        return ()


def _heading_level(paragraph) -> int | None:
    """Resolve built-in/custom paragraph styles through their basedOn chain."""
    for name in _paragraph_style_names(paragraph):
        match = re.match(r'^Heading\s+([1-6])$', name, re.I)
        if match:
            return int(match.group(1))
    return None


def _paragraph_has_style(paragraph, *names: str) -> bool:
    wanted = {name.casefold() for name in names}
    return any(name.casefold() in wanted for name in _paragraph_style_names(paragraph))


def _run_is_code(run, paragraph) -> bool:
    try:
        if any(str(style.name or '') == 'QuietWriter Code' for style in _style_chain(run.style)):
            return True
    except Exception:
        pass
    return _effective_font_name(run, paragraph) in {
        'consolas', 'courier new', 'courier', 'liberation mono'
    }


def _run_signature(run, paragraph) -> tuple[bool, bool, bool, bool, bool]:
    """Resolve run semantics while looking up Word styles only once per run.

    ``python-docx`` style lookup can scan the document style collection. The old
    implementation accessed ``run.style``/``paragraph.style`` repeatedly for
    code, font name and every boolean attribute. Large imported manuscripts paid
    that cost many times per run. Resolve both chains once and reuse them.
    """
    try:
        run_style = run.style
    except Exception:
        run_style = None
    try:
        paragraph_style = paragraph.style
    except Exception:
        paragraph_style = None
    run_chain = tuple(_style_chain(run_style))
    paragraph_chain = tuple(_style_chain(paragraph_style))

    def chain_value(chain, attr: str):
        for style in chain:
            try:
                value = getattr(style.font, attr)
            except Exception:
                value = None
            if value is not None:
                return value
        return None

    def effective_bool(direct_attr: str, *, on_font: bool = False) -> bool:
        try:
            direct = getattr(run.font if on_font else run, direct_attr)
        except Exception:
            direct = None
        if direct is not None:
            return bool(direct)
        value = chain_value(run_chain, direct_attr)
        if value is None:
            value = chain_value(paragraph_chain, direct_attr)
        return bool(value) if value is not None else False

    try:
        direct_font_name = run.font.name
    except Exception:
        direct_font_name = None
    font_name = direct_font_name or chain_value(run_chain, 'name') or chain_value(paragraph_chain, 'name') or ''
    code = any(str(style.name or '') == 'QuietWriter Code' for style in run_chain) or str(font_name).lower() in {
        'consolas', 'courier new', 'courier', 'liberation mono'
    }
    return (
        bool(code),
        effective_bool('bold'),
        effective_bool('italic'),
        effective_bool('underline'),
        effective_bool('strike', on_font=True),
    )


def _styles_for_signature(sig: tuple[bool, bool, bool, bool, bool]) -> frozenset[str]:
    code, bold, italic, underline, strike = sig
    if code:
        # Code is a literal inline region; other character emphasis has no useful
        # separate meaning inside it in QuietWriter's current syntax.
        return frozenset({'code'})
    return frozenset(
        name for name, enabled in (
            ('bold', bold), ('italic', italic), ('underline', underline), ('strike', strike)
        ) if enabled
    )


def _is_inside_deleted_text(element) -> bool:
    parent = element.getparent()
    while parent is not None:
        if parent.tag == qn('w:del'):
            return True
        parent = parent.getparent()
    return False


def _paragraph_runs_including_hyperlinks(paragraph):
    """Yield visible descendant runs in document order.

    Walking OOXML descendant runs instead of only ``paragraph.runs`` preserves
    hyperlink text, field results, tracked insertions, inline content controls
    and text-box text anchored in a paragraph. Deleted tracked text stays out.
    """
    try:
        for element in paragraph._p.iter():
            if element.tag != qn('w:r') or _is_inside_deleted_text(element):
                continue
            yield Run(element, paragraph)
        return
    except Exception:
        pass
    yield from paragraph.runs


def _run_import_text(run, warnings: list[str] | None = None) -> str:
    """Return writer-visible run text while retaining Word hard/soft breaks."""
    out: list[str] = []
    saw_content = False
    try:
        for child in run._r.iterchildren():
            if child.tag == qn('w:t'):
                out.append(child.text or '')
                saw_content = True
            elif child.tag == qn('w:tab'):
                out.append('\t')
                saw_content = True
            elif child.tag in {qn('w:br'), qn('w:cr')}:
                break_type = child.get(qn('w:type')) if child.tag == qn('w:br') else None
                out.append('\u2028')
                saw_content = True
                if break_type == 'page' and warnings is not None:
                    message = tr(
                        'docx.import.inline_page_break_as_line_break',
                        'Een pagina-einde midden in een alinea is als regeleinde geïmporteerd; paginalay-out is niet behouden.'
                    )
                    if message not in warnings:
                        warnings.append(message)
            elif child.tag == qn('w:noBreakHyphen'):
                out.append('-')
                saw_content = True
            elif child.tag == qn('w:softHyphen'):
                out.append('\u00ad')
                saw_content = True
    except Exception:
        pass
    if saw_content:
        return ''.join(out)
    return str(getattr(run, 'text', '') or '').replace('\n', '\u2028')


def paragraph_to_import_block(paragraph, warnings: list[str] | None = None) -> ImportBlock:
    """Read one Word paragraph into neutral import semantics.

    No QuietWriter markup punctuation is created here. That is intentionally
    deferred to the central import serializer.
    """
    chunks: list[tuple[tuple[bool, bool, bool, bool, bool], str]] = []
    for run in _paragraph_runs_including_hyperlinks(paragraph):
        text = _run_import_text(run, warnings)
        if not text:
            continue
        sig = _run_signature(run, paragraph)
        if chunks and chunks[-1][0] == sig:
            chunks[-1] = (sig, chunks[-1][1] + text)
        else:
            chunks.append((sig, text))

    style = _style_name(paragraph).lower()
    if style == 'quietwriter scene break':
        return ImportBlock(kind='scene')
    kind = 'paragraph'
    if style.startswith('list bullet'):
        kind = 'bullet'
    elif style.startswith('list number'):
        kind = 'numbered'
    elif style in {'quote', 'intense quote'}:
        kind = 'quote'

    runs = tuple(
        ImportInlineRun(text=text, styles=_styles_for_signature(sig))
        for sig, text in chunks
    )
    return ImportBlock(kind=kind, runs=runs)




def _docx_image_blocks(paragraph, assets: dict[str, ImportAsset], warnings: list[str]) -> list[ImportBlock]:
    """Extract supported Word drawings from one paragraph as neutral image blocks.

    QuietWriter currently models manuscript images as block elements. Word can
    anchor drawings inline with prose; when that occurs the visible paragraph
    text is kept and the image follows it as a block, with an explicit warning.
    """
    blocks: list[ImportBlock] = []
    try:
        blips = paragraph._p.xpath('.//a:blip')
    except Exception:
        return blocks
    if not blips:
        return blocks

    if (paragraph.text or '').strip():
        message = tr(
            'docx.import.inline_images_as_blocks',
            'Een of meer afbeeldingen stonden midden in een Word-alinea en zijn als losse afbeeldingsblokken geïmporteerd.'
        )
        if message not in warnings:
            warnings.append(message)

    for blip in blips:
        rel_id = blip.get(qn('r:embed'))
        if not rel_id:
            continue
        try:
            image_part = paragraph.part.related_parts[rel_id]
            data = bytes(image_part.blob)
            media_type = str(getattr(image_part, 'content_type', '') or '')
            filename = Path(str(getattr(image_part, 'partname', '') or 'image')).name or 'image'
        except Exception:
            continue
        try:
            decoded = DocxImage.from_blob(data)
            actual_media_type = str(decoded.content_type or '')
        except Exception:
            message = tr(
                'docx.import.corrupt_image_skipped',
                'Een beschadigde Word-afbeelding is overgeslagen; de overige tekst en afbeeldingen zijn wel geïmporteerd.'
            )
            if message not in warnings:
                warnings.append(message)
            continue
        if actual_media_type not in {'image/png', 'image/jpeg'}:
            message = tr(
                'docx.import.unsupported_image_type',
                'Een Word-afbeelding met type {type} is niet geïmporteerd; QuietWriter ondersteunt bij import PNG en JPEG.',
                type=actual_media_type or media_type or 'onbekend',
            )
            if message not in warnings:
                warnings.append(message)
            continue
        media_type = actual_media_type
        canonical_ext = '.png' if media_type == 'image/png' else '.jpg'
        if Path(filename).suffix.lower() not in {'.png', '.jpg', '.jpeg'}:
            filename = Path(filename).stem + canonical_ext

        digest = hashlib.sha256(data).hexdigest()
        key = f'image:{digest}'
        assets.setdefault(key, ImportAsset(key=key, filename=filename, media_type=media_type, data=data))

        alt = ''
        try:
            node = blip
            while node is not None and not node.tag.endswith(('}inline', '}anchor')):
                node = node.getparent()
            if node is not None:
                doc_pr = node.find(qn('wp:docPr'))
                if doc_pr is not None:
                    alt = str(doc_pr.get('descr') or doc_pr.get('title') or '')
        except Exception:
            alt = ''
        blocks.append(ImportBlock(kind='image', asset_key=key, alt=alt))
    return blocks


def _paragraph_to_import_blocks(paragraph, assets: dict[str, ImportAsset], warnings: list[str]) -> list[ImportBlock]:
    """Return paragraph semantics followed by any Word drawings it contains."""
    text_block = paragraph_to_import_block(paragraph, warnings)
    images = _docx_image_blocks(paragraph, assets, warnings)
    if text_block.visible_text or text_block.kind == 'scene' or not images:
        return [text_block, *images]
    return images


def paragraph_to_markdown(paragraph) -> str:
    """Compatibility adapter: serialize neutral Word semantics to source."""
    return serialize_import_block(paragraph_to_import_block(paragraph))



def _table_to_import_blocks(table: Table) -> list[ImportBlock]:
    """Flatten a Word table to readable row paragraphs without dropping text."""
    blocks: list[ImportBlock] = []
    for row in table.rows:
        cells: list[str] = []
        previous_cell = None
        for cell in row.cells:
            # python-docx repeats the same cell object for horizontally merged
            # cells; do not duplicate its visible text.
            identity = id(cell._tc)
            if identity == previous_cell:
                continue
            previous_cell = identity
            text = ' '.join(part.strip() for part in cell.text.splitlines() if part.strip())
            cells.append(text)
        row_text = ' · '.join(cells).strip()
        if row_text:
            blocks.append(ImportBlock(runs=(ImportInlineRun(row_text),)))
    return blocks


def _sdt_element_blocks(element) -> list[ImportBlock]:
    """Recover visible paragraphs from one block content control element."""
    blocks: list[ImportBlock] = []
    for paragraph in element.iter(qn('w:p')):
        text = ''.join(node.text or '' for node in paragraph.iter(qn('w:t'))).strip()
        if text:
            blocks.append(ImportBlock(runs=(ImportInlineRun(text),)))
    return blocks


def _top_level_sdt_blocks(doc) -> list[ImportBlock]:
    """Recover visible text from top-level block content controls."""
    blocks: list[ImportBlock] = []
    try:
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:sdt'):
                blocks.extend(_sdt_element_blocks(child))
    except Exception:
        return blocks
    return blocks


def _paragraph_has_page_break(paragraph) -> bool:
    try:
        for br in paragraph._p.xpath('.//w:br'):
            if br.get(qn('w:type')) == 'page':
                return True
        return bool(paragraph._p.xpath('./w:pPr/w:pageBreakBefore'))
    except Exception:
        return False


def _page_chunks(paragraphs) -> list[list]:
    pages: list[list] = [[]]
    for paragraph in paragraphs:
        pages[-1].append(paragraph)
        if _paragraph_has_page_break(paragraph):
            pages.append([])
    return [
        page for page in pages
        if any((p.text or '').strip() or bool(p._p.xpath('.//a:blip')) for p in page)
    ]


def _generic_word_title(value: str) -> bool:
    return value.strip().lower() in {'', 'document', 'word document', 'microsoft word document'}


def _display_heading(paragraph) -> bool:
    text = (paragraph.text or '').strip()
    if not text or len(text) > 140:
        return False
    if _paragraph_has_style(paragraph, 'Title', 'Subtitle') or _heading_level(paragraph) is not None:
        return True
    sizes = [run.font.size.pt for run in paragraph.runs if run.text.strip() and run.font.size]
    return bool(sizes and max(sizes) >= 18)


def _page_to_import_chapters(page, assets: dict[str, ImportAsset], warnings: list[str], chapter_offset: int = 0) -> list[ImportChapter]:
    chapters: list[ImportChapter] = []
    split_internal = len(page) > 25
    title_parts: list[str] = []
    body: list[ImportBlock] = []

    def flush():
        nonlocal title_parts, body
        if not title_parts and not any(block.visible_text.strip() or block.kind == 'image' for block in body):
            body = []
            return
        if not title_parts:
            first_block = next((block for block in body if block.visible_text.strip()), None)
            first = first_block.visible_text.strip() if first_block else 'Hoofdstuk'
            title_parts = [first[:97].rstrip() + ('…' if len(first) > 97 else '')]
            if first_block is not None:
                removed = False
                kept: list[ImportBlock] = []
                for block in body:
                    if not removed and block is first_block:
                        removed = True
                        continue
                    kept.append(block)
                body = kept
        chapters.append(ImportChapter(title=' — '.join(title_parts), blocks=tuple(body)))
        title_parts, body = [], []

    for paragraph in page:
        text = (paragraph.text or '').strip()
        if _display_heading(paragraph):
            if split_internal and body and any(block.visible_text.strip() or block.kind == 'image' for block in body):
                flush()
            elif body and any(block.visible_text.strip() or block.kind == 'image' for block in body):
                body.extend(_paragraph_to_import_blocks(paragraph, assets, warnings))
                continue
            if text:
                title_parts.append(text.replace('\n', ' '))
            continue
        paragraph_blocks = _paragraph_to_import_blocks(paragraph, assets, warnings)
        block = paragraph_blocks[0] if paragraph_blocks else ImportBlock()
        if not title_parts and not body and text:
            number = chapter_offset + len(chapters) + 1
            title_parts = [tr('docx.import.fallback_part', 'Deel {number}', number=number)]
            body.extend(paragraph_blocks)
            continue
        body.extend(paragraph_blocks)
    flush()
    return chapters


def _parse_page_oriented_docx(doc, source: Path, paragraphs, warnings: list[str], assets: dict[str, ImportAsset]) -> ImportDocument:
    pages = _page_chunks(paragraphs)
    chapters: list[ImportChapter] = []
    for page in pages:
        chapters.extend(_page_to_import_chapters(page, assets, warnings, chapter_offset=len(chapters)))

    core_title = str(doc.core_properties.title or '').strip()
    if _generic_word_title(core_title) and pages:
        first_nonempty = next(((p.text or '').strip() for p in pages[0] if (p.text or '').strip()), '')
        core_title = first_nonempty or source.stem
    return ImportDocument(
        title=core_title or source.stem,
        author=str(doc.core_properties.author or '').strip(),
        language=_document_default_language(doc),
        sections=(ImportSection(None, tuple(chapters)),),
        warnings=tuple(warnings),
        assets=tuple(assets.values()),
    )

def read_docx_import_document(source: Path) -> ImportDocument:
    """Read a DOCX conservatively into QuietWriter sections/chapters.

    QuietWriter-created DOCX files use custom paragraph styles for unambiguous
    section/chapter boundaries. For ordinary Word documents, Heading 1 is treated
    as chapter headings unless Heading 2 is also present; with Heading 2 present,
    Heading 1 becomes a section and Heading 2 a chapter.
    """
    source = Path(source)
    doc = Document(str(source))
    warnings: list[str] = []
    assets: dict[str, ImportAsset] = {}
    if doc.tables:
        warnings.append(tr('docx.import.tables_as_text', '{count} tabel(len) zijn als leesbare tekst geïmporteerd; tabelopmaak is niet behouden.', count=len(doc.tables)))
    # OOXML constructs python-docx can omit from paragraph.text/runs are detected
    # explicitly so imported content is never silently claimed to be complete.
    xml = doc.element
    try:
        has_revisions = bool(xml.xpath('.//w:ins | .//w:del'))
    except Exception:
        has_revisions = False
    if has_revisions:
        warnings.append(tr('docx.import.revisions_not_imported', 'Wijzigingen bijhouden is vereenvoudigd geïmporteerd: zichtbare ingevoegde tekst blijft behouden, verwijderde tekst en wijzigingsinformatie niet.'))
    try:
        has_controls = bool(xml.xpath('.//w:sdt'))
    except Exception:
        has_controls = False
    if has_controls:
        warnings.append(tr('docx.import.content_controls_not_imported', 'Tekst in Word-inhoudsbesturingselementen kan niet volledig zijn geïmporteerd.'))
    try:
        has_textboxes = bool(xml.xpath('.//w:txbxContent'))
    except Exception:
        has_textboxes = False
    if has_textboxes:
        warnings.append(tr('docx.import.textboxes_as_text', 'Tekst uit Word-tekstvakken is als gewone tekst geïmporteerd; positie en vormgeving zijn niet behouden.'))
    try:
        has_fields = bool(xml.xpath('.//w:fldSimple | .//w:instrText'))
    except Exception:
        has_fields = False
    if has_fields:
        warnings.append(tr('docx.import.fields_not_imported', 'Zichtbare resultaten van Word-velden zijn als gewone tekst geïmporteerd waar mogelijk; veldlogica is niet behouden.'))
    try:
        has_hyperlinks = any(rel.reltype == RT.HYPERLINK for rel in doc.part.rels.values())
    except Exception:
        has_hyperlinks = False
    if has_hyperlinks:
        warnings.append(tr('docx.import.hyperlink_targets_not_imported', 'Hyperlinktekst is geïmporteerd, maar hyperlinkadressen worden nog niet behouden.'))
    try:
        has_footnotes = bool(xml.xpath('.//w:footnoteReference'))
        has_endnotes = bool(xml.xpath('.//w:endnoteReference'))
    except Exception:
        has_footnotes = False
        has_endnotes = False
    if has_footnotes:
        warnings.append(tr('docx.import.footnotes_not_imported', 'Voetnoten worden nog niet geïmporteerd.'))
    if has_endnotes:
        warnings.append(tr('docx.import.endnotes_not_imported', 'Eindnoten worden nog niet geïmporteerd.'))

    paragraphs = list(doc.paragraphs)
    styles = [_style_name(p) for p in paragraphs]
    heading_levels = [_heading_level(p) for p in paragraphs]
    has_qw = any(s in {'QuietWriter Section', 'QuietWriter Chapter'} for s in styles)
    has_h2 = any(level == 2 for level in heading_levels)
    page_break_count = sum(1 for p in paragraphs if _paragraph_has_page_break(p))
    first_heading = next((i for i, level in enumerate(heading_levels) if level is not None), None)
    breaks_before_heading = 0 if first_heading is None else sum(1 for p in paragraphs[:first_heading] if _paragraph_has_page_break(p))
    has_title_style = any(_paragraph_has_style(p, 'Title', 'Subtitle') for p in paragraphs)
    if not doc.tables and not has_qw and page_break_count >= 2 and (has_title_style or breaks_before_heading >= 1):
        return _parse_page_oriented_docx(doc, source, paragraphs, warnings, assets)

    groups: list[tuple[str | None, list[ImportChapter]]] = []
    current_section: str | None = None
    current_title: str | None = None
    current_blocks: list[ImportBlock] = []
    preamble: list[ImportBlock] = []

    def ensure_group(title: str | None):
        nonlocal current_section
        if not groups or groups[-1][0] != title:
            groups.append((title, []))
        current_section = title

    def flush_chapter():
        nonlocal current_title, current_blocks, preamble
        if current_title is None:
            return
        if not groups:
            ensure_group(current_section)
        groups[-1][1].append(ImportChapter(current_title, tuple(current_blocks)))
        current_title = None
        current_blocks = []

    def start_chapter(title: str):
        nonlocal current_title, current_blocks, preamble
        flush_chapter()
        if not groups:
            ensure_group(current_section)
        current_title = title.strip() or 'Hoofdstuk'
        current_blocks = []
        if preamble:
            current_blocks.extend(preamble)
            preamble = []

    if hasattr(doc, 'iter_inner_content'):
        wrapped_items = list(doc.iter_inner_content())
        by_element = {getattr(item, '_element', None): item for item in wrapped_items}
        content_items = []
        for child in doc.element.body.iterchildren():
            wrapped = by_element.get(child)
            if wrapped is not None:
                content_items.append(wrapped)
            elif child.tag == qn('w:sdt'):
                blocks = _sdt_element_blocks(child)
                if blocks:
                    content_items.append(blocks)
    else:
        content_items = paragraphs
    for item in content_items:
        if isinstance(item, list):
            if current_title is None:
                preamble.extend(item)
            else:
                current_blocks.extend(item)
            continue
        if isinstance(item, Table):
            table_blocks = _table_to_import_blocks(item)
            if current_title is None:
                preamble.extend(table_blocks)
            else:
                current_blocks.extend(table_blocks)
            continue
        p = item
        style = _style_name(p)
        text = (p.text or '').strip()
        heading_level = _heading_level(p)
        is_section = style == 'QuietWriter Section' or (not has_qw and has_h2 and heading_level == 1)
        is_chapter = style == 'QuietWriter Chapter' or (
            not has_qw and ((has_h2 and heading_level == 2) or (not has_h2 and heading_level == 1))
        )
        if is_section and text:
            flush_chapter()
            ensure_group(text)
            continue
        if is_chapter and text:
            start_chapter(text)
            continue
        paragraph_blocks = _paragraph_to_import_blocks(p, assets, warnings)
        if current_title is None:
            # QuietWriter DOCX files contain a generated title/front-matter area
            # before the first explicit chapter. Those publication pages are not
            # manuscript chapters on re-import. External DOCX files keep their
            # pre-heading text as a first chapter.
            if not has_qw:
                preamble.extend(paragraph_blocks)
        else:
            current_blocks.extend(paragraph_blocks)

    flush_chapter()
    if preamble:
        if not groups:
            groups.append((None, []))
        title = str(doc.core_properties.title or '').strip() or source.stem
        groups[-1][1].insert(0, ImportChapter(title, tuple(preamble)))
    if not any(chapters for _, chapters in groups):
        title = str(doc.core_properties.title or '').strip() or source.stem
        groups = [(None, [ImportChapter(title, ())])]

    return ImportDocument(
        title=str(doc.core_properties.title or '').strip() or source.stem,
        author=str(doc.core_properties.author or '').strip(),
        language=_document_default_language(doc),
        sections=tuple(ImportSection(section, tuple(chapters)) for section, chapters in groups),
        warnings=tuple(warnings),
        assets=tuple(assets.values()),
    )


def parse_docx_book(source: Path) -> DocxImportResult:
    """Compatibility boundary used by the current storage workflow.

    The DOCX reader itself returns neutral import semantics. Only here are those
    semantics serialized to QuietWriter source.
    """
    imported = read_docx_import_document(source)
    warnings = list(imported.warnings)
    if imported.assets:
        warnings.append(tr(
            'docx.import.legacy_result_images_omitted',
            'Afbeeldingen zijn gelezen maar worden alleen bij volledige boekimport in de mediastore geplaatst.'
        ))
    return DocxImportResult(
        title=imported.title, author=imported.author, language=imported.language,
        sections=serialize_import_document(imported), warnings=tuple(warnings),
    )


def _ensure_styles(doc: Document):
    styles = doc.styles
    if 'QuietWriter Section' not in styles:
        style = styles.add_style('QuietWriter Section', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles['Heading 1']
    if 'QuietWriter Chapter' not in styles:
        style = styles.add_style('QuietWriter Chapter', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles['Heading 2']
    if 'QuietWriter Code' not in styles:
        style = styles.add_style('QuietWriter Code', WD_STYLE_TYPE.CHARACTER)
        style.font.name = 'Consolas'
        style.font.size = Pt(10)
    if 'QuietWriter Scene Break' not in styles:
        style = styles.add_style('QuietWriter Scene Break', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles['Normal']


def _append_inline(paragraph, source: str, doc: Document, runs=None):
    """Append central semantic inline runs to a DOCX paragraph."""
    semantic_runs = tuple(runs) if runs is not None else inline_runs(source)
    if not semantic_runs:
        if source:
            paragraph.add_run(source)
        return
    for semantic in semantic_runs:
        _add_run(paragraph, semantic.text, semantic.styles, doc)


def _add_run(paragraph, text: str, styles: frozenset[str], doc: Document):
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold = 'bold' in styles
    run.italic = 'italic' in styles
    run.underline = 'underline' in styles
    run.font.strike = 'strike' in styles
    if 'code' in styles:
        run.style = doc.styles['QuietWriter Code']



def _language_tag(language: str) -> str:
    value = (language or '').strip().replace('_', '-').lower()
    if value.startswith('nl'):
        return 'nl-NL'
    if value.startswith('en'):
        return 'en-US'
    if value.startswith('de'):
        return 'de-DE'
    if value.startswith('fr'):
        return 'fr-FR'
    if value.startswith('es'):
        return 'es-ES'
    return language or 'nl-NL'


def _set_rpr_language(rpr, tag: str):
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), tag)
    lang.set(qn('w:eastAsia'), tag)


def _set_document_language(doc: Document, language: str):
    tag = _language_tag(language)
    styles_el = doc.styles.element
    defaults = styles_el.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults')
        styles_el.insert(0, defaults)
    rpr_default = defaults.find(qn('w:rPrDefault'))
    if rpr_default is None:
        rpr_default = OxmlElement('w:rPrDefault'); defaults.append(rpr_default)
    rpr = rpr_default.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rpr_default.append(rpr)
    _set_rpr_language(rpr, tag)
    for style_name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Quote', 'List Bullet', 'List Number', 'QuietWriter Section', 'QuietWriter Chapter', 'QuietWriter Scene Break'):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        rpr = style.element.get_or_add_rPr()
        _set_rpr_language(rpr, tag)


def _content_box(section):
    return (
        int(section.page_width - section.left_margin - section.right_margin),
        max(1, int(section.page_height - section.top_margin - section.bottom_margin - Pt(18))),
    )


def _add_picture_fit(paragraph, data: bytes, max_width: int, max_height: int, *, allow_upscale: bool = False):
    image = DocxImage.from_blob(data)
    width, height = int(image.width), int(image.height)
    if width <= 0 or height <= 0:
        return paragraph.add_run().add_picture(BytesIO(data))
    scale = min(max_width / width, max_height / height)
    if not allow_upscale:
        scale = min(scale, 1.0)
    width = max(1, int(width * scale)); height = max(1, int(height * scale))
    return paragraph.add_run().add_picture(BytesIO(data), width=width, height=height)

def _asset_map(export_document) -> dict[str, object]:
    result = {}
    for asset in export_document.assets:
        if asset.role != 'inline':
            continue
        result[asset.filename] = asset
        result[Path(asset.filename).name] = asset
    return result


def _append_markdown(doc: Document, markdown: str, export_document):
    assets = _asset_map(export_document)
    source = (markdown or '').replace('\r\n', '\n').replace('\r', '\n')
    for block in parse_document(source).blocks:
        if block.kind == 'empty':
            doc.add_paragraph()
            continue
        if block.kind == 'image':
            attrs = block.attrs or {}
            image_path = str(attrs.get('path') or '')
            if not image_path:
                continue
            image_alt = str(attrs.get('alt') or '')
            image_caption = str(attrs.get('caption') or '')
            p = doc.add_paragraph()
            asset = assets.get(Path(image_path).name) or assets.get(image_path)
            if asset:
                try:
                    max_width, max_height = _content_box(doc.sections[-1])
                    _add_picture_fit(p, asset.data, max_width, max_height)
                except Exception:
                    _append_inline(p, image_alt or Path(image_path).name, doc)
            else:
                _append_inline(p, image_alt or Path(image_path).name, doc)
            if image_caption:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cp.add_run(image_caption); r.italic = True
            continue

        if block.kind == 'scene':
            p = doc.add_paragraph(serialize_block_source('scene', ''), style='QuietWriter Scene Break')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        content = block.text[block.content_start-block.start:block.content_end-block.start]
        runs = content_inline_runs(block)
        if block.kind == 'heading':
            p = doc.add_paragraph(style='Heading 3')
            _append_inline(p, content, doc, runs)
            continue
        if block.kind == 'quote':
            p = doc.add_paragraph(style='Quote')
            _append_inline(p, content, doc, runs)
            continue
        if block.kind == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
            _append_inline(p, content, doc, runs)
            continue
        if block.kind == 'numbered':
            p = doc.add_paragraph(style='List Number')
            _append_inline(p, content, doc, runs)
            continue

        p = doc.add_paragraph()
        _append_inline(p, content, doc, runs)


def export_docx(export_document, destination: Path, settings: dict | None = None) -> Path:
    """Export a QuietWriter snapshot to an editable DOCX document."""
    destination = Path(destination)
    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    _ensure_styles(doc)
    _set_document_language(doc, export_document.language)
    doc.core_properties.title = export_document.title
    doc.core_properties.author = export_document.author

    if export_document.cover_asset:
        cover_p = doc.add_paragraph()
        cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        max_width, max_height = _content_box(doc.sections[-1])
        try:
            _add_picture_fit(cover_p, export_document.cover_asset.data, max_width, max_height, allow_upscale=True)
        except Exception:
            pass

    title = doc.add_paragraph(style='Title')
    if export_document.cover_asset:
        title.paragraph_format.page_break_before = True
    title.add_run(export_document.title)
    if export_document.author:
        author = doc.add_paragraph()
        author.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author.add_run(export_document.author)
    doc.add_page_break()

    # Textual publication front matter remains editable Word text. Structured
    # title/copyright widgets are deliberately not duplicated beyond the title page.
    for item in export_document.front_matter:
        if item.kind == 'text' and item.text.strip():
            _append_markdown(doc, item.text, export_document)
            doc.add_page_break()

    multiple_sections = len(export_document.sections) > 1
    started_chapter = False
    for section in export_document.sections:
        show_section = multiple_sections or (section.title and section.id != 'root' and section.title.lower() != 'manuscript')
        if show_section:
            section_p = doc.add_paragraph(section.title, style='QuietWriter Section')
            if started_chapter:
                section_p.paragraph_format.page_break_before = True
        for chapter_index, chapter in enumerate(section.chapters):
            chapter_p = doc.add_paragraph(chapter.title, style='QuietWriter Chapter')
            if started_chapter and not (show_section and chapter_index == 0):
                chapter_p.paragraph_format.page_break_before = True
            _append_markdown(doc, chapter.markdown, export_document)
            started_chapter = True

    for item in export_document.back_matter:
        if item.kind == 'text' and item.text.strip():
            doc.add_page_break()
            _append_markdown(doc, item.text, export_document)

    destination.parent.mkdir(parents=True, exist_ok=True)
    tmp = destination.with_suffix(destination.suffix + '.tmp')
    doc.save(str(tmp))
    tmp.replace(destination)
    return destination
