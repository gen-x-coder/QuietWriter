"""Exportdoel "Inzenden": een standaard manuscript voor uitgeverij of wedstrijd.

Wat een redacteur of jury verwacht, en wat hier dus gebeurt:
- A4, marges 2,5 cm, één lettertype (standaard Times New Roman 12);
- dubbele regelafstand (of 1,5), geen witregels tussen alinea's;
- alinea's springen in, behalve de eerste na een hoofdstuktitel of scènebreuk;
- een scènebreuk is een gecentreerde ``#``;
- elk hoofdstuk begint op een nieuwe pagina, een stuk onder de bovenrand;
- een titelpagina met je naam en contactgegevens (of anoniem, voor wedstrijden), het geschatte aantal woorden,
  de titel en je naam; geen omslag, colofon of opdracht;
- een kopregel "Achternaam / TITEL / paginanummer" vanaf de eerste tekstpagina;
- "EINDE" na het laatste hoofdstuk.

De teksten op de titelpagina volgen de **boektaal**, niet de taal van de app.
Afbeeldingen worden als korte aanduiding opgenomen: een inzending is tekst.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..document_view import content_inline_runs, parse_document
from ..manuscript_text import count_words
from ..docx_io import _add_run, _append_inline, _ensure_styles, _set_document_language
from ..document_view import InlineRun

FONTS = {
    'times': 'Times New Roman',
    'arial': 'Arial',
    'courier': 'Courier New',
}
SPACINGS = {'double': WD_LINE_SPACING.DOUBLE, 'one_half': WD_LINE_SPACING.ONE_POINT_FIVE}

DEFAULT_SUBMISSION = {
    'font': 'times',
    'line_spacing': 'double',
    'running_header': True,
    'anonymous': False,
}

_WORDS = {
    'nl': ('door', 'ongeveer {n} woorden', 'EINDE', 'Afbeelding'),
    'en': ('by', 'about {n} words', 'THE END', 'Image'),
    'de': ('von', 'etwa {n} Wörter', 'ENDE', 'Abbildung'),
    'fr': ('par', 'environ {n} mots', 'FIN', 'Image'),
    'es': ('por', 'unas {n} palabras', 'FIN', 'Imagen'),
}
_THOUSANDS = {'nl': '.', 'de': '.', 'es': '.', 'fr': ' ', 'en': ','}


def _lang(language: str) -> str:
    code = (language or 'nl').strip().lower()[:2]
    return code if code in _WORDS else 'en'


def submission_options(settings: dict | None) -> dict:
    options = dict(DEFAULT_SUBMISSION)
    options.update((settings or {}).get('submission') or {})
    if options.get('font') not in FONTS:
        options['font'] = 'times'
    if options.get('line_spacing') not in SPACINGS:
        options['line_spacing'] = 'double'
    options['running_header'] = bool(options.get('running_header', True))
    options['anonymous'] = bool(options.get('anonymous', False))
    options['contact'] = str(options.get('contact') or '').strip()
    return options


def rounded_word_count(words: int) -> int:
    """Zoals uitgeverijen het willen: afgerond, niet tot op het woord."""
    if words < 1000:
        return max(0, int(round(words / 10.0)) * 10)
    if words < 20000:
        return int(round(words / 100.0)) * 100
    return int(round(words / 1000.0)) * 1000


def format_number(value: int, language: str) -> str:
    sep = _THOUSANDS.get(_lang(language), ',')
    return f'{value:,}'.replace(',', sep)


def surname(author: str) -> str:
    """'Lucas Bonsel' → 'Bonsel'; 'Jan van der Berg' → 'van der Berg'."""
    parts = (author or '').split()
    return ' '.join(parts[1:]) if len(parts) > 1 else (parts[0] if parts else '')


def short_title(title: str) -> str:
    return (title or '').split(':')[0].strip() or (title or '').strip()


def manuscript_word_count(document) -> int:
    return sum(count_words(ch.markdown) for section in document.sections for ch in section.chapters)


def _add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = f' {instruction} '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t'); value.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for element in (begin, text, sep, value, end):
        run._r.append(element)


def _start_page_numbers_at(section, number: int):
    sect_pr = section._sectPr
    pg = sect_pr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sect_pr.append(pg)
    pg.set(qn('w:start'), str(number))


class _Writer:
    def __init__(self, doc: Document, options: dict, language: str):
        self.doc = doc
        self.options = options
        self.labels = _WORDS[_lang(language)]
        self.indent_next = False

    def paragraph(self, *, indent: bool | None = None, align=None, style=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0); fmt.space_after = Pt(0)
        fmt.line_spacing_rule = SPACINGS[self.options['line_spacing']]
        use_indent = self.indent_next if indent is None else indent
        fmt.first_line_indent = Cm(1.25) if use_indent else Cm(0)
        if align is not None:
            p.alignment = align
        return p

    def chapter_title(self, title: str, *, new_page: bool):
        p = self.paragraph(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.page_break_before = new_page
        p.paragraph_format.space_before = Pt(144)   # ongeveer een derde van de pagina
        p.paragraph_format.space_after = Pt(24)
        p.paragraph_format.keep_with_next = True
        p.add_run(title)
        self.indent_next = False

    def scene_break(self):
        self.paragraph(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER).add_run('#')
        self.indent_next = False

    def markdown(self, source: str):
        source = (source or '').replace('\r\n', '\n').replace('\r', '\n')
        for block in parse_document(source).blocks:
            if block.kind == 'empty':
                continue
            if block.kind == 'scene':
                self.scene_break()
                continue
            if block.kind == 'image':
                attrs = block.attrs or {}
                label = str(attrs.get('caption') or attrs.get('alt') or Path(str(attrs.get('path') or '')).name)
                p = self.paragraph(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
                p.add_run(f'[{self.labels[3]}: {label}]' if label else f'[{self.labels[3]}]').italic = True
                self.indent_next = False
                continue
            content = block.text[block.content_start - block.start:block.content_end - block.start]
            runs = content_inline_runs(block)
            if block.kind == 'heading':
                p = self.paragraph(indent=False)
                p.paragraph_format.keep_with_next = True
                if runs:
                    _append_inline(p, content, self.doc,
                                   [InlineRun(r.text, frozenset(r.styles) | {'bold'}, r.start, r.end) for r in runs])
                if not runs and content:
                    _add_run(p, content, frozenset({'bold'}), self.doc)
                self.indent_next = False
                continue
            if block.kind == 'quote':
                p = self.paragraph(indent=False)
                p.paragraph_format.left_indent = Cm(1.25)
            elif block.kind in ('bullet', 'numbered'):
                p = self.paragraph(indent=False)
                p.paragraph_format.left_indent = Cm(1.25)
                prefix = '– ' if block.kind == 'bullet' else ''
                if prefix:
                    p.add_run(prefix)
            else:
                p = self.paragraph()
            if runs:
                _append_inline(p, content, self.doc, runs)   # met echte voetnoten (1.7)
            if not runs and content:
                _add_run(p, content, frozenset(), self.doc)
            self.indent_next = True


def export_submission(document, destination: Path, settings: dict | None = None) -> Path:
    options = submission_options(settings)
    destination = Path(destination)
    language = document.language or 'nl'
    labels = _WORDS[_lang(language)]
    doc = Document()
    _ensure_styles(doc)
    _set_document_language(doc, language)
    font_name = FONTS[options['font']]
    normal = doc.styles['Normal']
    normal.font.name = font_name
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(attr), font_name)
    doc.core_properties.title = document.title
    doc.core_properties.author = '' if options['anonymous'] else (document.author or '')
    doc.core_properties.last_modified_by = doc.core_properties.author
    doc.core_properties.comments = ''

    title_section = doc.sections[0]
    for section in (title_section,):
        section.page_width = Cm(21.0); section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)

    writer = _Writer(doc, options, language)

    # --- titelpagina (eigen sectie, zonder kopregel) ---------------------------------
    anonymous = options['anonymous']
    author = '' if anonymous else (document.author or '')
    contact_lines = [] if anonymous else ([line.rstrip() for line in options['contact'].splitlines()] or [author])
    for line in contact_lines:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    words = rounded_word_count(manuscript_word_count(document))
    p = doc.add_paragraph(labels[1].format(n=format_number(words, language)))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(200)
    p.add_run((document.title or '').upper()).bold = True
    if author:
        p = doc.add_paragraph(f'{labels[0]} {author}')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)

    # --- tekst (nieuwe sectie: kopregel en paginanummer vanaf 1) ------------------------
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.page_width = Cm(21.0); body.page_height = Cm(29.7)
    body.left_margin = body.right_margin = Cm(2.5)
    body.top_margin = body.bottom_margin = Cm(2.5)
    body.header.is_linked_to_previous = False
    _start_page_numbers_at(body, 1)
    if options['running_header']:
        header = body.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        parts = [part for part in (surname(author), short_title(document.title).upper()) if part]
        header.add_run(' / '.join(parts) + ' / ')
        _add_field(header, 'PAGE')
    title_section.header.is_linked_to_previous = False

    multiple_sections = len(document.sections) > 1
    first = True
    for section in document.sections:
        if multiple_sections and section.title:
            writer.chapter_title(section.title.upper(), new_page=not first)
            first = False
        for chapter in section.chapters:
            writer.chapter_title(chapter.title, new_page=not first)
            first = False
            writer.markdown(chapter.markdown)
    end = writer.paragraph(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    end.paragraph_format.space_before = Pt(24)
    end.add_run(labels[2])

    destination.parent.mkdir(parents=True, exist_ok=True)
    tmp = destination.with_suffix(destination.suffix + '.tmp')
    doc.save(str(tmp))
    tmp.replace(destination)
    return destination
