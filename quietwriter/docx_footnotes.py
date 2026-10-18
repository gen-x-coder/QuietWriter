"""Echte Word-voetnoten voor DOCX-export (1.7).

python-docx kent geen voetnoten. Dit kleine hulpje maakt het onderdeel
``word/footnotes.xml`` aan (met de twee scheidingsnoten die Word verwacht),
koppelt het aan het document en voegt per noot een verwijzing in de tekst toe.
Word en LibreOffice nummeren de noten zelf en zetten ze onder aan de pagina.
"""
from __future__ import annotations

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

FOOTNOTES_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'

_EMPTY_FOOTNOTES = (
    f'<w:footnotes {nsdecls("w")}>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '</w:footnotes>'
)

# Elementen in w:settings die volgens het schema ná w:footnotePr komen.
_SETTINGS_AFTER_FOOTNOTE_PR = (
    'w:endnotePr', 'w:compat', 'w:docVars', 'w:rsids', 'm:mathPr', 'w:attachedSchema', 'w:themeFontLang',
    'w:clrSchemeMapping', 'w:doNotIncludeSubdocsInStats', 'w:doNotAutoCompressPictures', 'w:forceUpgrade',
    'w:captions', 'w:readModeInkLockDown', 'w:smartTagType', 'sl:schemaLibrary', 'w:shapeDefaults',
    'w:doNotEmbedSmartTags', 'w:decimalSymbol', 'w:listSeparator',
)


class DocxFootnotes:
    """Eén per document: ``DocxFootnotes.for_document(doc)``."""

    def __init__(self, doc):
        self.doc = doc
        self.root = None
        self.next_id = 1

    @classmethod
    def for_document(cls, doc) -> 'DocxFootnotes':
        existing = getattr(doc, '_quietwriter_footnotes', None)
        if existing is None:
            existing = cls(doc)
            doc._quietwriter_footnotes = existing
        return existing

    def _ensure_part(self):
        if self.root is not None:
            return
        document_part = self.doc.part
        for rel in document_part.rels.values():
            if rel.reltype == RT.FOOTNOTES:
                self.root = rel.target_part.element
                ids = [int(el.get(qn('w:id'))) for el in self.root.findall(qn('w:footnote'))]
                self.next_id = max([0, *ids]) + 1
                return
        element = parse_xml(_EMPTY_FOOTNOTES)
        part = XmlPart(PackURI('/word/footnotes.xml'), FOOTNOTES_CONTENT_TYPE, element, document_part.package)
        document_part.relate_to(part, RT.FOOTNOTES)
        self.root = element
        self._register_in_settings()

    def _register_in_settings(self):
        settings = self.doc.settings.element
        if settings.find(qn('w:footnotePr')) is not None:
            return
        pr = OxmlElement('w:footnotePr')
        for value in ('-1', '0'):
            ref = OxmlElement('w:footnote'); ref.set(qn('w:id'), value); pr.append(ref)
        later = None
        for child in settings:
            tag = child.tag
            if any(tag == qn(name) for name in _SETTINGS_AFTER_FOOTNOTE_PR if not name.startswith(('m:', 'sl:'))) \
                    or tag.endswith('}mathPr') or tag.endswith('}schemaLibrary'):
                later = child
                break
        if later is not None:
            later.addprevious(pr)
        else:
            settings.append(pr)

    def add(self, paragraph, note_runs, add_run) -> int:
        """Zet een verwijzing aan het eind van *paragraph* en maak de noot.

        *add_run(paragraph, text, styles)* schrijft één opgemaakte run; zo krijgt
        de noottekst dezelfde opmaak (vet, cursief …) als in de editor.
        """
        self._ensure_part()
        note_id = self.next_id
        self.next_id += 1

        ref_run = paragraph.add_run()
        ref_run.font.superscript = True
        ref = OxmlElement('w:footnoteReference'); ref.set(qn('w:id'), str(note_id))
        ref_run._r.append(ref)

        footnote = OxmlElement('w:footnote'); footnote.set(qn('w:id'), str(note_id))
        p = OxmlElement('w:p')
        ppr = OxmlElement('w:pPr')
        spacing = OxmlElement('w:spacing'); spacing.set(qn('w:after'), '0'); ppr.append(spacing)
        p.append(ppr)
        footnote.append(p)
        self.root.append(footnote)
        note_paragraph = Paragraph(p, paragraph._parent)
        mark = note_paragraph.add_run()
        mark.font.superscript = True
        mark._r.append(OxmlElement('w:footnoteRef'))
        note_paragraph.add_run(' ')
        for run in note_runs:
            add_run(note_paragraph, run.text, run.styles)
        for run in note_paragraph.runs:
            run.font.size = Pt(10)   # python-docx houdt de schemavolgorde in w:rPr aan
        return note_id
